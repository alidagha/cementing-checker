"""Synthetic structural fixtures only; no App-generated golden or engineering PASS."""
from io import BytesIO
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import zipfile

from docx import Document
from docx.oxml import OxmlElement
import pytest

from checker_parser import ROOT, parse_docx


def save(doc, path):
    doc.save(path)
    return path


def table(doc, title, rows):
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).merge(t.cell(0, 1)).text = title
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    return t


def body(result):
    assert result['status'] == 'EXTRACTED_UNVERIFIED', result
    return result['records']


@pytest.mark.parametrize('slurries', [('Main',), ('Lead #1', 'Lead #2', 'Tail')])
def test_repeated_lab_groups_and_explicit_table_association(tmp_path, slurries):
    doc = Document()
    doc.add_paragraph('VII. Laboratory Report')
    for i, fluid in enumerate(slurries):
        doc.add_paragraph(f'{fluid} Cement Slurry Results')
        table(doc, 'Test Basic Data', [('BHCT', str(100 + i))])
        table(doc, 'Conventional', [('Material Type', 'Mass (g/600 mL)'), ('Cement', str(500 + i))])
    records = body(parse_docx(save(doc, tmp_path / 'groups.docx')))
    for i, fluid in enumerate(slurries):
        rec = next(r for r in records if r['raw_text'] == str(500 + i))
        assert rec['association'] == fluid
        assert rec['column_label'] == 'Mass (g/600 mL)'
        assert rec['locator']['row'] == 2
    assert len([r for r in records if r['raw_text'] == 'Test Basic Data']) == len(slurries)


def test_cement_table_title_split_runs_and_duplicate_labels(tmp_path):
    doc = Document()
    doc.add_paragraph('V. Cementing Program')
    for fluid in ('Lead', 'Tail'):
        t = table(doc, '', [('Volume', '40 bbl'), ('Volume', '50 bbl')])
        p = t.cell(0, 0).paragraphs[0]
        for part in (fluid, ' Cement ', 'Slurry Data'):
            p.add_run(part)
    records = body(parse_docx(save(doc, tmp_path / 'split.docx')))
    volumes = [r for r in records if r['raw_text'] in ('40 bbl', '50 bbl')]
    assert [r['association'] for r in volumes] == ['Lead', 'Lead', 'Tail', 'Tail']
    assert len({json.dumps(r['locator'], sort_keys=True) for r in volumes}) == 4


def test_ambiguity_never_silently_overridden(tmp_path):
    doc = Document()
    doc.add_paragraph('VII. Laboratory Report')
    doc.add_paragraph('Lead Cement Slurry Results')
    table(doc, 'Tail Cement Slurry Results', [('BHCT', '100')])
    records = body(parse_docx(save(doc, tmp_path / 'conflict.docx')))
    rec = next(r for r in records if r['raw_text'] == '100')
    assert rec['association_state'] == 'AMBIGUOUS'
    assert rec['association'] is None
    assert rec['association_candidates'] == ['Lead', 'Tail']
    assert all(r['verification_status'] == 'NOT_VERIFIABLE' for r in records)


def test_unheaded_lab_is_not_attached_by_position(tmp_path):
    doc = Document()
    doc.add_paragraph('VII. Laboratory Report')
    table(doc, 'Test Basic Data', [('BHCT', '100')])
    rec = next(r for r in body(parse_docx(save(doc, tmp_path / 'unheaded.docx'))) if r['raw_text'] == '100')
    assert rec['association_state'] == 'NOT_VERIFIABLE'


def test_missing_zero_na_dash_and_malformed_distinct(tmp_path):
    doc = Document()
    doc.add_paragraph('V. Cementing Program')
    table(doc, 'Main Cement Slurry Data', [('Volume', x) for x in ('', '0', 'N/A', '-', 'oops')])
    records = body(parse_docx(save(doc, tmp_path / 'states.docx')))
    states = [r['value_state'] for r in records if r['printed_label'] == 'Volume' and r['locator'].get('column') == 1]
    assert states == ['MISSING', 'ZERO', 'NA_PRINTED', 'DASH', 'MALFORMED']
    assert all(r['value_state'] == 'TEXT' for r in records if r['raw_text'] == 'Volume')


def test_full_ordered_prose_notes_and_nested_table(tmp_path):
    doc = Document()
    doc.add_paragraph('Executive Summary')
    p = doc.add_paragraph()
    for run in ('Meaningful ', 'summary & ', 'details.'):
        p.add_run(run)
    doc.add_paragraph('NOTE 1: summary water')
    doc.add_paragraph('X. Cementing Procedure')
    doc.add_paragraph('1. Prepare water.\nKeep all lines.')
    outer = table(doc, 'Instructions', [('Step detail', 'Preserve this')])
    inner = outer.cell(1, 1).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = 'Nested instruction'
    doc.add_paragraph('NOTE 9: Do not skip.')
    doc.add_paragraph('2. Pump cement.')
    records = body(parse_docx(save(doc, tmp_path / 'prose.docx')))
    texts = [r['raw_text'] for r in records]
    expected = ['Meaningful summary & details.', 'NOTE 1: summary water', '1. Prepare water.\nKeep all lines.', 'Nested instruction', 'NOTE 9: Do not skip.', '2. Pump cement.']
    assert [t for t in texts if t in expected] == expected
    assert [r['step'] for r in records if r['step']] == [1, 2]
    assert [r['note'] for r in records if r['note']] == [1, 9]


def test_exclusions_and_optional_absent_fluids(tmp_path):
    doc = Document()
    doc.add_paragraph('Cover secret')
    doc.add_paragraph('I. Document Control')
    doc.add_paragraph('Approval secret')
    doc.add_paragraph('IV. Fluids Sequence')
    table(doc, 'Fluids', [('Type', 'Density'), ('Main', '104 pcf')])
    doc.add_paragraph('NOTE 6: The Pore and Frac Pressure is Based on Mud Weight Density.')
    doc.add_paragraph('VII. Laboratory Report')
    table(doc, 'Rheology Test', [('IOD', 'excluded value'), ('RPM', '300')])
    doc.add_paragraph('VIII. Directional Well Survey Data')
    doc.add_paragraph('Survey secret')
    doc.add_paragraph('IX. Centralizer Placement Recommendation')
    doc.add_paragraph('NOTE 20: excluded centralizer')
    doc.add_paragraph('X. Cementing Procedure')
    doc.add_paragraph('1. Pump Main.')
    records = body(parse_docx(save(doc, tmp_path / 'exclude.docx')))
    text = '\n'.join(r['raw_text'] for r in records)
    assert 'secret' not in text and 'excluded' not in text and 'IOD' not in text
    assert 'Pore and Frac' not in text
    assert not any(r['association'] in ('Spacer', 'Pre Flush', 'Diesel') for r in records)
    assert next(r for r in records if r['raw_text'] == '104 pcf')['association'] == 'Main'


def test_deterministic_original_unchanged_and_cli(tmp_path):
    doc = Document()
    doc.add_paragraph('Executive Summary')
    doc.add_paragraph('Summary claims.')
    path = save(doc, tmp_path / 'deterministic.docx')
    before = path.read_bytes()
    first = parse_docx(path)
    assert first == parse_docx(path)
    assert first['source_document_sha256'] == hashlib.sha256(before).hexdigest()
    proc = subprocess.run([sys.executable, '-B', str(ROOT / 'checker_parser.py'), str(path)], capture_output=True, text=True)
    assert proc.returncode == 0
    assert json.loads(proc.stdout) == first
    assert path.read_bytes() == before


@pytest.mark.parametrize('kind', ['junk', 'xml', 'traversal', 'dtd', 'unsupported', 'external', 'extension'])
def test_damaged_unsupported_and_unsafe_fail_closed(tmp_path, kind):
    path = tmp_path / ('bad.txt' if kind == 'extension' else 'bad.docx')
    if kind in ('junk', 'extension'):
        path.write_bytes(b'not Word')
    else:
        doc = Document()
        doc.add_paragraph('Executive Summary')
        if kind == 'unsupported':
            doc.element.body.append(OxmlElement('w:sdt'))
        buf = BytesIO()
        doc.save(buf)
        with zipfile.ZipFile(buf) as src, zipfile.ZipFile(path, 'w') as dst:
            for name in src.namelist():
                raw = src.read(name)
                if name == 'word/document.xml':
                    if kind == 'xml': raw = b'<broken'
                    if kind == 'dtd': raw = b'<!DOCTYPE x [<!ENTITY y SYSTEM "file:///etc/passwd">]>' + raw
                if kind == 'external' and name == 'word/_rels/document.xml.rels':
                    raw = raw.replace(b'</Relationships>', b'<Relationship Id="evil" Type="x" TargetMode="External" Target="https://example.com/"/></Relationships>')
                dst.writestr(name, raw)
            if kind == 'traversal': dst.writestr('../escaped', 'never extract')
    before = path.read_bytes()
    result = parse_docx(path)
    assert result['status'] == 'REJECTED', result
    assert result['records'] == []
    assert path.read_bytes() == before


def test_symlink_escape_and_traversal_input(tmp_path):
    link = tmp_path / 'escape.docx'
    link.symlink_to('/workspace/cementing-test/master_template.docx')
    assert parse_docx(link)['status'] == 'REJECTED'
    assert parse_docx(ROOT / '../cementing-test/master_template.docx')['status'] == 'REJECTED'


def test_parser_does_not_write_or_contact_network(tmp_path):
    doc = Document()
    doc.add_paragraph('Executive Summary')
    doc.add_paragraph('No side effects.')
    path = save(doc, tmp_path / 'readonly.docx')
    # Audit in a fresh process, after imports, before parse. Deny all writes/network.
    code = '''import sys
from checker_parser import parse_docx
def audit(event, args):
    if event == 'open':
        mode, flags = args[1:3]
        if (isinstance(mode, str) and any(c in mode for c in 'wax+')) or (isinstance(flags, int) and flags & 3):
            raise AssertionError(('write', args))
    if event.startswith(('socket.', 'os.mkdir', 'os.remove', 'os.rename', 'os.symlink')):
        raise AssertionError((event, args))
sys.addaudithook(audit)
assert parse_docx(sys.argv[1])['status'] == 'EXTRACTED_UNVERIFIED'
'''
    proc = subprocess.run([sys.executable, '-B', '-c', code, str(path)], cwd=ROOT, capture_output=True, text=True)
    assert proc.returncode == 0, proc.stderr


def test_explicit_stages_and_styled_lab_heading(tmp_path):
    doc = Document()
    doc.add_paragraph('VII. Laboratory Report')
    doc.add_heading('Main Cement Slurry Results', level=2)
    table(doc, 'Test Basic Data', [('BHCT', '100')])
    doc.add_paragraph('X. Cementing Procedure')
    doc.add_paragraph('Stage 2:')
    doc.add_paragraph('1. Pump water.')
    doc.add_paragraph('Stage 3:')
    doc.add_paragraph('2. Pump Main.')
    records = body(parse_docx(save(doc, tmp_path / 'stages.docx')))
    bhct = next(r for r in records if r['raw_text'] == '100')
    assert bhct['section'] == 'lab' and bhct['association'] == 'Main'
    assert [r['stage'] for r in records if r['step']] == [2, 3]


def test_unknown_body_wrapper_fails_closed(tmp_path):
    doc = Document()
    doc.add_paragraph('Executive Summary')
    unknown = OxmlElement('w:customXml')
    para = OxmlElement('w:p')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = 'Do not silently lose this summary'
    run.append(text)
    para.append(run)
    unknown.append(para)
    doc.element.body.append(unknown)
    assert parse_docx(save(doc, tmp_path / 'unknown.docx'))['status'] == 'REJECTED'


# C01R: pinned-template exclusions and blank-heading context recovery.
def template_wrappers():
    from docx.oxml.ns import qn
    doc = Document(ROOT / 'snapshot/master_template.docx')
    body = doc.element.body
    toc = next(body.iter(qn('w:sdt')))
    boxes = list(body.iter(qn('w:txbxContent')))
    assert len(boxes) == 4
    cover = boxes[0]
    control = boxes[-1]
    while cover.getparent() is not body:
        cover = cover.getparent()
    while control.getparent() is not body:
        control = control.getparent()
    return deepcopy(cover), deepcopy(toc), deepcopy(control)


def test_c01r_excluded_pinned_wrappers_preserve_program_and_order(tmp_path):
    from docx.oxml.ns import qn
    cover, toc, control = template_wrappers()
    doc = Document()
    doc.element.body.insert(0, cover)
    doc.add_paragraph('Executive Summary')
    doc.add_paragraph('Full summary with significant water preparation wording.')
    doc.element.body.insert(len(doc.element.body)-1, toc)
    doc.add_paragraph('I. Document Control')
    doc.element.body.insert(len(doc.element.body)-1, control)
    doc.add_paragraph('Approval-only content')
    doc.add_paragraph('V. Cementing Program')
    table(doc, 'Tail Cement Slurry Data', [('Volume', '40 bbl')])
    doc.add_paragraph('NOTE 7: Keep the complete mixing instruction.')
    doc.add_paragraph('VII. Laboratory Report')
    doc.add_paragraph('Tail Cement Slurry Results')
    table(doc, 'Test Basic Data', [('BHCT', '100 degF')])
    doc.add_paragraph('VIII. Directional Well Survey Data')
    doc.add_paragraph('Excluded survey content')
    doc.add_heading('', level=2)
    doc.add_paragraph('Still excluded survey content')
    doc.add_paragraph('IX. Centralizer Placement Recommendation')
    doc.add_paragraph('Excluded centralizer content')
    doc.add_paragraph('X. Cementing Procedure')
    doc.add_paragraph('1. Prepare water.\nPreserve the second line.')
    doc.add_paragraph('NOTE 8: Keep the second instruction.')
    doc.add_paragraph('2. Pump Tail.')
    path = save(doc, tmp_path / 'pinned-wrappers.docx')
    before = path.read_bytes()
    result = parse_docx(path)
    records = body(result)
    texts = [r['raw_text'] for r in records]
    expected = ['Full summary with significant water preparation wording.',
                'NOTE 7: Keep the complete mixing instruction.',
                '1. Prepare water.\nPreserve the second line.',
                'NOTE 8: Keep the second instruction.', '2. Pump Tail.']
    assert [s for s in texts if s in expected] == expected
    assert not any('excluded' in s.casefold() or 'Approval-only' in s for s in texts)
    assert next(r for r in records if r['raw_text'] == '40 bbl')['association'] == 'Tail'
    assert next(r for r in records if r['raw_text'] == '100 degF')['section'] == 'lab'
    assert len(result['excluded_wrappers']) == 5
    assert all(x['xml_path'] and x['body_block'] is not None for x in result['excluded_wrappers'])
    assert parse_docx(path) == result
    cli = subprocess.run([sys.executable, '-B', str(ROOT / 'checker_parser.py'), str(path)], capture_output=True, text=True)
    assert cli.returncode == 0 and json.loads(cli.stdout) == result
    assert path.read_bytes() == before


@pytest.mark.parametrize('blank', ['', '  \t\n '])
def test_c01r_blank_heading_preserves_known_section_and_lab_identity(tmp_path, blank):
    doc = Document()
    doc.add_paragraph('V. Cementing Program')
    doc.add_heading(blank, level=2)
    table(doc, 'Lead #1 Cement Slurry Data', [('Volume', '40 bbl')])
    doc.add_paragraph('VII. Laboratory Report')
    doc.add_heading('Lead #1 Cement Slurry Results', level=2)
    doc.add_heading(blank, level=2)
    table(doc, 'Test Basic Data', [('BHCT', '100 degF')])
    records = body(parse_docx(save(doc, tmp_path / 'blank.docx')))
    for text, section in [('40 bbl', 'cement'), ('100 degF', 'lab')]:
        row = next(r for r in records if r['raw_text'] == text)
        assert row['section'] == section and row['association'] == 'Lead #1'
    assert next(r for r in records if r['raw_text'] == '40 bbl')['heading'] == 'V. Cementing Program'


@pytest.mark.parametrize('kind', ['sdt', 'textbox', 'forged_toc', 'control_nonempty'])
def test_c01r_included_or_unproven_wrappers_rejected(tmp_path, kind):
    from docx.oxml.ns import qn
    _, toc, control = template_wrappers()
    doc = Document()
    doc.add_paragraph('Executive Summary')
    doc.add_paragraph('Keep included summary')
    if kind == 'sdt':
        wrapper = OxmlElement('w:sdt')
        content = OxmlElement('w:sdtContent')
        content.append(deepcopy(doc.paragraphs[-1]._p))
        wrapper.append(content)
    elif kind in ('textbox', 'control_nonempty'):
        wrapper = control
        for box in wrapper.iter(qn('w:txbxContent')):
            box.append(deepcopy(doc.paragraphs[-1]._p))
        if kind == 'control_nonempty':
            doc.add_paragraph('I. Document Control')
    else:
        wrapper = toc
        wrapper.find(qn('w:sdtContent')).append(deepcopy(doc.paragraphs[-1]._p))
    doc.element.body.insert(len(doc.element.body)-1, wrapper)
    path = save(doc, tmp_path / 'included-wrapper.docx')
    before = path.read_bytes()
    result = parse_docx(path)
    assert result['status'] == 'REJECTED' and result['records'] == []
    assert path.read_bytes() == before


def test_c01r_true_unknown_heading_resets_scope_and_identity(tmp_path):
    doc = Document()
    doc.add_paragraph('VII. Laboratory Report')
    doc.add_paragraph('Lead Cement Slurry Results')
    doc.add_heading('Unrelated appendix', level=1)
    table(doc, 'Test Basic Data', [('BHCT', '101')])
    doc.add_paragraph('X. Cementing Procedure')
    doc.add_paragraph('1. Resume known procedure.')
    records = body(parse_docx(save(doc, tmp_path / 'nonblank.docx')))
    unknown = next(r for r in records if r['raw_text'] == '101')
    assert unknown['section'] == 'unknown' and unknown['association'] is None
    assert next(r for r in records if r['step'] == 1)['section'] == 'procedure'


def test_c01r_pinned_wrappers_parse_without_writes_or_network():
    code = '''import sys
from checker_parser import parse_docx
def audit(event,args):
    if event == 'open':
        mode,flags=args[1:3]
        if (isinstance(mode,str) and any(c in mode for c in 'wax+')) or (isinstance(flags,int) and flags & (3|64|512|1024)):
            raise AssertionError(('write',args))
    if event.startswith(('socket.','os.mkdir','os.remove','os.rename','os.symlink')):
        raise AssertionError((event,args))
sys.addaudithook(audit)
result=parse_docx('snapshot/master_template.docx')
assert result['status']=='EXTRACTED_UNVERIFIED',result
assert len(result['excluded_wrappers'])==5
'''
    proc = subprocess.run([sys.executable, '-B', '-c', code], cwd=ROOT, capture_output=True, text=True)
    assert proc.returncode == 0, proc.stderr
