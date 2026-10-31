"""Freshly bound, read-only ordinary-slurry DV impact through accepted C04."""
import argparse
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn
from c10_corrections import ROOT, Refusal, digest, identities, literal
from checker_parser import safe_input, text_of, heading
from c02_engine import read_json

sys.path.insert(0, str(ROOT/'c04'))
import c04_runtime as runtime
import c04_preview as preview_source

SCHEMA = 'g06b.1'


def private_path(value, suffix):
    path = Path(value).absolute()
    if (path.resolve(strict=True) != path or path.suffix != suffix or not path.is_file()
            or not path.is_relative_to(runtime.HERE/'tmp')):
        raise Refusal('PRIVATE_C04_INPUT_WITHOUT_SYMLINK_REQUIRED')
    if any((p/'.git').exists() for p in (path.parent, *path.parents)):
        raise Refusal('GIT_FREE_RUNTIME_REQUIRED')
    return runtime.local_path(path)


def plain_paragraph(element, document, raw):
    literal(element, document)
    if element.tag != qn('w:p') or any(n.tag not in (qn('w:pPr'), qn('w:r'), qn('w:proofErr')) for n in element):
        raise Refusal('PLAIN_PHYSICAL_PARAGRAPH_REQUIRED')
    # The pinned template includes spelling/grammar markers, which carry no text.
    for marker in element.findall(qn('w:proofErr')):
        if (len(marker) or marker.text or set(marker.attrib) != {qn('w:type')}
                or marker.get(qn('w:type')) not in ('spellStart', 'spellEnd', 'gramStart', 'gramEnd')):
            raise Refusal('NONLITERAL_PROOFING_MARKER')
    runs = element.findall(qn('w:r'))
    if len(runs) != 1 or any(n.tag not in (qn('w:rPr'), qn('w:t')) for n in runs[0]):
        raise Refusal('SPLIT_OR_COMPLEX_CLAIM')
    texts = runs[0].findall(qn('w:t'))
    if len(texts) != 1 or texts[0].text != raw or text_of(element) != raw:
        raise Refusal('EXACT_PHYSICAL_CLAIM_REQUIRED')
    if set(texts[0].attrib) - {'{http://www.w3.org/XML/1998/namespace}space'}:
        raise Refusal('UNSUPPORTED_TEXT_ATTRIBUTES')


def claim_proof(original, baseline, slurry):
    water = [w for w in baseline['water_findings'] if w['slurry'] == slurry]
    if len(water) != 1 or water[0]['status'] != 'INFERRED_AT_DISPLAY_PRECISION':
        raise Refusal('UNIQUE_C03_NOTE_TANK_ASSOCIATION_REQUIRED')
    water = water[0]
    document = Document(BytesIO(original)); body = document.element.body
    sections = [p for p in body if p.tag == qn('w:p') and heading(text_of(p)) == 'cement']
    if len(sections) != 1:
        raise Refusal('UNIQUE_CEMENT_SECTION_REQUIRED')
    literal(sections[0], document)
    for field in ('yield', 'sacks'):
        matches = [f for f in baseline['c02']['findings'] if f['slurry'] == slurry and f['field'] == field]
        if (len(matches) != 1 or matches[0]['mapping_state'] != 'IDENTIFIED'
                or matches[0]['association_state'] != 'EXPLICIT' or matches[0]['section'] != 'cement'):
            raise Refusal('UNIQUE_ORDINARY_SLURRY_IDENTITY_REQUIRED')
    tables = water['proof_table_locators']
    if len(tables) != 3 or len({p['body_block'] for p in tables}) != 3:
        raise Refusal('UNIQUE_C03_TABLE_GROUP_REQUIRED')
    for loc, suffix in zip(tables, ('Data', 'Blend Data', 'Additives Data')):
        if loc['table_path']:
            raise Refusal('NESTED_SLURRY_GROUP')
        table = body[loc['body_block']]
        if table.tag != qn('w:tbl'):
            raise Refusal('PHYSICAL_SLURRY_TABLE_REQUIRED')
        literal(table, document)
        rows = table.findall(qn('w:tr'))
        title = rows[0].findall(qn('w:tc')) if rows else []
        paragraphs = title[0].findall(qn('w:p')) if len(title) == 1 else []
        if len(paragraphs) != 1:
            raise Refusal('UNIQUE_PHYSICAL_SLURRY_TITLE_REQUIRED')
        plain_paragraph(paragraphs[0], document, slurry+' Cement Slurry '+suffix)
    note = body[water['locator']['body_block']]
    if water['locator'].get('table') is not None:
        raise Refusal('PHYSICAL_PREPARATION_NOTE_REQUIRED')
    plain_paragraph(note, document, water['raw_text'])
    gross = [g for g in baseline['gross_additive_findings'] if g['slurry'] == slurry]
    expected = baseline['c02']['reference']['slurries'][slurry]['additives']
    if (len(gross) != len(expected) or any(g['status'] == 'NOT_VERIFIABLE' for g in gross)
            or sorted((g['material'], g['material_type']) for g in gross)
            != sorted((r['Name'], r['Material Type']) for r in expected)):
        raise Refusal('COMPLETE_UNIQUE_GROSS_ROW_ASSOCIATION_REQUIRED')
    table = body[tables[2]['body_block']]
    # The locked source has eight grid columns but seven physical additive cells.
    # Prove its actual grid against the template; never repair/guess that structure.
    template = Document(ROOT/'snapshot/master_template.docx')
    canonical = [t._tbl for t in template.tables
        if t.cell(0, 0).text == '{{ s.name }} Cement Slurry Additives Data']
    grid = table.find(qn('w:tblGrid'))
    if len(canonical) != 1 or grid is None or grid.xml != canonical[0].find(qn('w:tblGrid')).xml:
        raise Refusal('SOURCE_ADDITIVE_GRID_NOT_PROVED')
    physical = table.findall(qn('w:tr'))
    if len(physical) != len(expected) + 2:
        raise Refusal('MISSING_OR_EXTRA_PHYSICAL_GROSS_ROW')
    source_header = canonical[0].findall(qn('w:tr'))[1]
    header_cells = source_header.findall(qn('w:tc'))
    source_after = source_header.findall('./'+qn('w:trPr')+'/'+qn('w:gridAfter'))
    source_width = source_header.findall('./'+qn('w:trPr')+'/'+qn('w:wAfter'))
    if (len(header_cells) != 7 or len(grid) != 8 or len(source_after) != 1
            or source_after[0].attrib != {qn('w:val'): '1'} or len(source_width) != 1):
        raise Refusal('PINNED_TRAILING_GRID_LAYOUT_NOT_PROVED')
    if [''.join(text_of(c).split()) for c in physical[1].findall(qn('w:tc'))] != [
            ''.join(text_of(c).split()) for c in header_cells]:
        raise Refusal('SOURCE_GROSS_HEADERS_NOT_PROVED')
    for row in physical[1:]:
        cells = row.findall(qn('w:tc'))
        if len(cells) != 7 or any(n.tag in {qn('w:'+x) for x in
                ('gridSpan', 'vMerge', 'hMerge', 'gridBefore', 'wBefore', 'tbl')} for n in row.iter()):
            raise Refusal('MERGED_OR_IRREGULAR_GROSS_ROW')
        # Only the pinned trailing omission is admitted. The seven unspanned
        # physical cells occupy columns 0..6; the omitted column is column 7.
        for tag, reference in (('gridAfter', source_after[0]), ('wAfter', source_width[0])):
            markers = row.findall('./'+qn('w:trPr')+'/'+qn('w:'+tag))
            if (len(markers) != 1 or list(row.iter(qn('w:'+tag))) != markers
                    or markers[0].attrib != reference.attrib or len(markers[0])
                    or markers[0].text or (markers[0].tail or '').strip()):
                raise Refusal('UNPROVED_TRAILING_GRID_OMISSION')
        if len(cells) + int(source_after[0].get(qn('w:val'))) != len(grid):
            raise Refusal('GROSS_ROW_GRID_COVERAGE_MISMATCH')
    seen = set()
    for finding in gross:
        loc = finding['locator']
        if (loc['body_block'] != tables[2]['body_block'] or loc['table'] != tables[2]['table']
                or loc['table_path'] or loc['column'] != 6 or loc['column_span'] != 1
                or loc['row'] not in range(2, len(physical)) or loc['row'] in seen):
            raise Refusal('C03_PHYSICAL_GROSS_LOCATOR_NOT_PROVED')
        cells = physical[loc['row']].findall(qn('w:tc'))
        if (text_of(cells[1]) != finding['material'] or text_of(cells[2]) != finding['material_type']
                or text_of(cells[6]) != finding['raw_text']):
            raise Refusal('C03_PHYSICAL_GROSS_CLAIM_NOT_PROVED')
        seen.add(loc['row'])
    for row in physical[2:]:
        for cell in row.findall(qn('w:tc')):
            paragraphs = cell.findall(qn('w:p'))
            if len(paragraphs) != 1 or any(n.tag not in (qn('w:tcPr'), qn('w:p')) for n in cell):
                raise Refusal('SPLIT_PHYSICAL_GROSS_CELL')
            plain_paragraph(paragraphs[0], document, text_of(paragraphs[0]))
    return {'water': water, 'gross': gross, 'table_locators': tables}


def unresolved():
    return [{'dependency': name, 'classification': classification} for name, classification in (
        ('NOTE_Procedure_Summary_and_tank_instructions', 'MANUAL_REVIEW_REQUIRED'),
        ('timing_and_rates', 'NOT_VERIFIABLE'), ('geometry_and_excess', 'NOT_VERIFIABLE'),
        ('measured_Lab_and_rheology', 'NOT_VERIFIABLE'), ('CT_and_Spacer_cascades', 'NOT_VERIFIABLE'))]


def execute(docx, operator, slurry=None, edit=None, json_inputs=None):
    if edit is not None:
        if (slurry is not None or type(edit) is not dict or set(edit) !=
                {'schema', 'bindings', 'slurry', 'corrected_dead_volume_bbl'} or edit['schema'] != SCHEMA):
            raise Refusal('EXACTLY_ONE_BOUND_DV_EDIT_REQUIRED')
        slurry = edit['slurry']
        preview_source.validate_edit({k: edit[k] for k in ('slurry', 'corrected_dead_volume_bbl')})
    if type(slurry) is not str or not slurry:
        raise Refusal('EXPLICIT_EXACT_SLURRY_REQUIRED')
    path = private_path(docx, '.docx'); original = safe_input(path)
    operator_sha = digest(operator); edit_sha = digest(edit); pins = identities()
    def verify_files():
        if json_inputs is None: return
        if type(json_inputs) is not list or len(json_inputs) != (2 if edit is not None else 1):
            raise Refusal('INVALID_JSON_BINDINGS')
        for item, value in zip(json_inputs, [operator, edit]):
            raw = private_path(item['path'], '.json').read_bytes()
            if sha256(raw).hexdigest() != item['sha256'] or read_json(raw) != value:
                raise Refusal('INPUT_JSON_DRIFT')
    verify_files()
    payload = {'docx': str(path), 'operator': operator}
    if edit is None:
        baseline = runtime.child(dict(payload, action='baseline')); c04 = None
    else:
        c04 = runtime.child(dict(payload, action='preview', edit={k: edit[k] for k in ('slurry', 'corrected_dead_volume_bbl')}))
        if c04['status'] != 'PARTIAL_FINDINGS_ONLY' or c04['selected_dead_volume_preview'] is None:
            raise Refusal('C04_SELECTED_PROJECTION_NOT_VERIFIABLE')
        baseline = c04['c03']
    # child() installs the unchanged C04 audit/path boundary before source execution.
    preview_source.bind(baseline, operator, sha256(original).hexdigest())
    import materials_db as db
    import project_state as state_source
    if any(not Path(m.__file__).resolve().is_relative_to(ROOT/'snapshot') for m in (db, state_source)):
        raise Refusal('NON_SNAPSHOT_ROLE_AUTHORITY')
    if operator['job_type'] in (db.CT_JOB, 'CMT SQUEEZE') or slurry not in db.PLACEMENT_SLURRIES:
        raise Refusal('UNSUPPORTED_JOB_OR_NONORDINARY_ROLE')
    if (slurry not in operator['active_fluids'] or slurry not in baseline['c02']['reference']['slurries']
            or not all(slurry in state_source.available_fluid_types({'job_type': operator['job_type'], 'mud_type': mud}) for mud in ('WBM', 'OBM'))):
        raise Refusal('SELECTED_ROLE_APPLICABILITY_UNPROVED')
    proof = claim_proof(original, baseline, slurry)
    bindings = dict(pins, original_document_sha256=sha256(original).hexdigest(),
        operator_sha256=operator_sha, selected_slurry=slurry, c03_sha256=digest(baseline),
        selected_claims_sha256=digest(proof))
    if edit is not None and edit['bindings'] != bindings:
        raise Refusal('STALE_OR_FORGED_DV_SELECTION')
    result = {'schema': SCHEMA, 'status': 'PARTIAL_FINDINGS_ONLY', 'bindings': bindings,
        'operation': 'SELECTION_INSPECTION' if edit is None else 'EXPLICIT_DV_IMPACT',
        'original_c03': baseline, 'selected_claims': proof, 'c04': c04,
        'docx_written': False, 'word_modified': False, 'unresolved_dependencies': unresolved()}
    if c04 is not None:
        p = c04['selected_dead_volume_preview']
        result['impact_ledger'] = {'before_qualification': 'INFERRED_NONUNIQUE_AT_WORD_DISPLAY_PRECISION',
            'before_dv_interval': p['original_inferred_dead_volume'],
            'explicit_dead_volume_bbl': p['proposed_dead_volume_bbl'],
            'historical_exact_delta': None, 'historical_change': 'NOT_VERIFIABLE',
            'pure_mix_water_full_precision_bbl': p['pure_mix_water_full_precision_bbl'],
            'source_after_total_preparation_water_full_precision_bbl': p['new_preparation_water_full_precision_bbl'],
            'source_after_NOTE_display_bbl': p['new_preparation_water_note_display_bbl'],
            'source_after_Word_numeric_projection': p['source_word_numeric_projection'],
            'NOTE_display_equals_original': p['new_preparation_water_note_display_bbl'] == proof['water']['Total Preparation Water Reported in Word']['raw'],
            'gross_additives': [{'material': g['material'], 'material_type': g['material_type'],
                'unit': g['unit'], 'original_Word_records': g['word_records'],
                'before_source_display_enclosure': g['before_source_display_range'],
                'source_after_display': g['after_source_display'], 'source_after_full_precision': None,
                'full_precision_status': 'NOT_VERIFIABLE: original table API returns display only',
                'historical_exact_delta': None} for g in p['gross_additive_preview']],
            'source_invariants_and_unselected_slurries': p['unaffected'],
            'source_after_tables': p['after_source_tables']}
    if (identities() != pins or safe_input(private_path(path, '.docx')) != original
            or digest(operator) != operator_sha or digest(edit) != edit_sha):
        raise Refusal('INPUT_OR_SOURCE_DRIFT')
    verify_files()
    return result


def call(payload):
    environment = dict(os.environ, PYTHONDONTWRITEBYTECODE='1', TMPDIR=str(runtime.HERE/'tmp'),
        XDG_CACHE_HOME=str(runtime.HERE/'cache'), XDG_CONFIG_HOME=str(runtime.HERE/'cache/config'),
        MPLCONFIGDIR=str(runtime.HERE/'cache/mpl'))
    process = subprocess.run([sys.executable, '-B', str(Path(__file__).resolve()), '--worker'],
        input=json.dumps(payload, allow_nan=False), text=True, capture_output=True, cwd=ROOT,
        env=environment, timeout=120)
    if process.returncode not in (0, 2): raise Refusal('C04_IMPACT_PROCESS_FAILED')
    result = read_json(process.stdout)
    if (process.returncode == 0) != (result.get('status') == 'PARTIAL_FINDINGS_ONLY'):
        raise Refusal('C04_IMPACT_PROCESS_STATUS_MISMATCH')
    return result


def inspect_selection(docx, operator, slurry):
    return call({'docx': str(docx), 'operator': operator, 'slurry': slurry})


def dead_volume_impact(docx, operator, edit):
    return call({'docx': str(docx), 'operator': operator, 'edit': edit})


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--worker', action='store_true'); parser.add_argument('operator_json', nargs='?')
    parser.add_argument('docx', nargs='?'); mode = parser.add_mutually_exclusive_group()
    mode.add_argument('--slurry'); mode.add_argument('--edit'); args = parser.parse_args()
    try:
        if args.worker:
            result = execute(**read_json(sys.stdin.read()))
        else:
            if not args.operator_json or not args.docx or (args.slurry is None and args.edit is None):
                raise Refusal('EXPLICIT_INPUTS_AND_SELECTION_REQUIRED')
            files = [private_path(args.operator_json, '.json')] + ([private_path(args.edit, '.json')] if args.edit else [])
            raw = [p.read_bytes() for p in files]
            inputs = [{'path': str(p), 'sha256': sha256(v).hexdigest()} for p, v in zip(files, raw)]
            result = call({'docx': args.docx, 'operator': read_json(raw[0]), 'slurry': args.slurry,
                'edit': read_json(raw[1]) if args.edit else None, 'json_inputs': inputs})
            if raw != [private_path(p, '.json').read_bytes() for p in files]:
                raise Refusal('CLI_INPUT_DRIFT')
    except (ValueError, TypeError, KeyError, IndexError, OSError, RuntimeError, ArithmeticError, zipfile.BadZipFile) as error:
        result = {'schema': SCHEMA, 'status': 'INPUT_REJECTED', 'docx_written': False,
            'word_modified': False, 'reason': str(error) if isinstance(error, Refusal) else 'C04_INPUT_OR_SOURCE_NOT_VERIFIABLE'}
    print(json.dumps(result, sort_keys=True, allow_nan=False))
    return 0 if result['status'] == 'PARTIAL_FINDINGS_ONLY' else 2


if __name__ == '__main__': raise SystemExit(main())
