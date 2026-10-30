"""C06 source-backed well/geometry controls, not fifteen-job golden certification."""
from pathlib import Path
import tempfile
from copy import deepcopy
import hashlib
import json
import subprocess
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

from comparison_records import ROOT
from engineering_source import request
from c06_well import compare_document


def hardware(description, md, size, inner):
    return dict(zip(('Description','MD (m)','Size (in)','ID (in)'),(description,md,size,inner)),
        **{'Joint (m)':12.2,'Weight (ppf)':50,'Grade':'K-55','Collapse (psi)':1000,'Burst (psi)':2000})


def operator():
    job='CSG 9 5/8"'
    return {'job_type':job, 'hardware':[hardware('Casing','1000','9.625',8.535),
        hardware('Previous Casing','500','13.375',12.25), hardware('Open Hole Size','500-1000','12.25',12.25)],
        'placement_config':{'job_type':job,'target_row':'__manual__','manual_depth_m':1000,'chart_excess_mode':'ignore'},
        'well_geometry':{'version':1,'type':'Vertical','td_m':1000},
        'geothermal_config':{'version':1,'source':'BHST','surface_temp':80,'value':180},
        'fluids':{'Main':{'volume':50,'density':118,'pump_rate':2}}}


@pytest.fixture(scope='module')
def case():
    with tempfile.TemporaryDirectory(dir=ROOT/'tmp',prefix='c06-') as folder:
        p=Path(folder)/'synthetic.docx';op=operator()
        source=request('render_well',op,docx=str(p))
        yield p,op,source


def test_well_hardware_and_independent_geometry_oracles(case):
    p,op,source=case
    result=compare_document(p,op);ref=result['reference']['reference']
    assert result['status']=='PARTIAL_FINDINGS_ONLY'
    direct=[f for f in result['findings'] if f['field'].startswith(('hardware.','well.')) and f['full_precision_expected'] is not None and f['source_rounded_expected']!='-']
    assert len(direct)==23 and all(f['result']=='MATCH' for f in direct)
    assert ref['well']['full_precision']['geo_tvd']==1000  # independent vertical TD
    assert ref['well']['full_precision']['bhst']==180     # independent measured source
    assert ref['well']['display']['geo_gradient']=='3.05'
    assert ref['geometry']['model']['shoe_md']==1000
    assert ref['geometry']['client_requested_volumes_bbl']=={'Main':50}
    assert ref['geometry']['full_annular_capacity_bbl']>50
    assert next(f for f in result['findings'] if f['field']=='geometry.reference')['result']=='INFORMATIONAL'
    assert 'cement_chart_model.py:resolve_cement_placement' in source['source_calls']
    assert 'well_profile.py:resolve_well_profile' in source['source_calls']


def test_source_matrix_has_all_fifteen_jobs_without_golden_claim(case):
    _,_,source=case;matrix=source['reference']['job_support_matrix']
    assert len(matrix)==15
    assert {j for j,s in matrix.items() if s['chart_geometry']=='NOT_VERIFIABLE'}=={'CMT PLUG','CMT SQUEEZE','CMT PLUG WITH CT'}
    assert all(s['golden_parity']=='NOT_RUN' for s in matrix.values())
    assert {s['chart_geometry'] for s in matrix.values()}=={'casing','liner','tieback','NOT_VERIFIABLE'}


def test_client_volume_is_not_geometry_and_two_slurries_remain_distinct(case):
    _,op,source=case;op=deepcopy(op)
    op['fluids']={'Lead':{'volume':40,'density':100,'pump_rate':2},'Tail':{'volume':60,'density':118,'pump_rate':2}}
    ref=request('well',op)['reference']['geometry']
    assert ref['status']=='INFORMATIONAL'
    assert ref['client_requested_volumes_bbl']=={'Lead':40,'Tail':60}
    assert ref['full_annular_capacity_bbl']==source['reference']['geometry']['full_annular_capacity_bbl']
    assert {r['entered_bbl'] for r in ref['placement']}=={40,60}


def test_apply_excess_changes_capacity_without_replacing_client_input(case):
    _,op,source=case;op=deepcopy(op)
    op['placement_config'].update(chart_excess_mode='apply',excess_csg_oh_pct=10,excess_csg_csg_pct=10)
    ref=request('well',op)['reference']['geometry']
    assert ref['full_annular_capacity_bbl']==pytest.approx(1.1*source['reference']['geometry']['full_annular_capacity_bbl'])
    assert ref['client_requested_volumes_bbl']=={'Main':50}


@pytest.mark.parametrize('kind',['liner','tieback'])
def test_source_exceptional_installed_paths(kind):
    op=operator();op.pop('well_geometry');op.pop('geothermal_config')
    if kind=='liner':
        op['job_type']='LNR 5"'
        op['hardware']=[hardware('Drill Pipes','500','5',4),hardware('Liner Hanger','500','5',4),
            hardware('Liner','500-1000','5',4),hardware('Previous Casing','600','9.625',8.535),hardware('Open Hole Size','600-1000','8.5',8.5)]
        op['placement_config'].update(job_type=op['job_type'])
    else:
        op['job_type']='TIE BACK LNR 7"'
        op['hardware']=[hardware('Tie Back','0-700','7',6),hardware('Previous Casing','800','9.625',8.535),
            hardware('Previous Liner','700-1200','7',6),hardware('Open Hole Size','800-1200','8.5',8.5)]
        op.pop('placement_config')
        choices=request('well',op)['reference']['source_hardware_choices']
        host=next(k for k,v in choices.items() if v['description']=='Previous Casing')
        op['placement_config']={'job_type':op['job_type'],'target_row':'__manual__','manual_depth_m':700,'chart_excess_mode':'ignore','host_row':host}
    ref=request('well',op)['reference']['geometry']
    assert ref['status']=='INFORMATIONAL' and ref['model']['kind']==kind


@pytest.mark.parametrize('job',['CMT PLUG','CMT SQUEEZE','CMT PLUG WITH CT'])
def test_unsupported_chart_jobs_do_not_borrow_casing_engine(job):
    op=operator();op['job_type']=job;op['placement_config']['job_type']=job
    if job.endswith('WITH CT'):op['hardware']=[hardware('CT',500,'0',0)]
    ref=request('well',op)['reference']
    assert ref['geometry']['status']=='NOT_VERIFIABLE'
    if job.endswith('WITH CT'):assert ref['ct_setting_depth_m']==500


@pytest.mark.parametrize('bad',['target','gap','diameter','mode','volume'])
def test_invalid_geometry_is_unverified_not_guessed(case,bad):
    _,op,_=case;op=deepcopy(op)
    if bad=='target':op['placement_config']['manual_depth_m']=None
    if bad=='gap':op['hardware'][2]['MD (m)']='600-1000'
    if bad=='diameter':op['hardware'][0]['ID (in)']=20
    if bad=='mode':op['placement_config']['chart_excess_mode']=None
    if bad=='volume':op['fluids']['Main']['volume']=0
    assert request('well',op)['reference']['geometry']['status']=='NOT_VERIFIABLE'


@pytest.mark.parametrize('mutation',['md','unit','duplicate','crossed','temperature','split'])
def test_untrusted_word_boundaries(case,monkeypatch,mutation,tmp_path):
    p,op,source=case
    # Mapping controls reuse this fixture's actual source calculation; no new physics.
    monkeypatch.setattr('c06_well.request',lambda *a,**k: source)
    d=Document(p);hw=next(t for t in d.tables if t.cell(0,0).text=='Tubular and Casing Hardware')
    casing=next(r for r in hw.rows if r.cells[0].text=='Casing')
    if mutation=='md':casing.cells[1].text='999.0'
    if mutation=='unit':next(r for r in hw.rows if r.cells[1].text=='m' and r.cells[2].text=='in').cells[1].text='ft'
    if mutation=='duplicate':hw._tbl.addnext(deepcopy(hw._tbl))
    if mutation=='crossed':casing.cells[0].text='Previous Casing'
    if mutation=='temperature':
        t=next(t for t in d.tables if t.cell(0,0).text=='Temperature');t._tbl.addnext(deepcopy(t._tbl))
    if mutation=='split':casing.cells[1].add_paragraph('1000.0')
    target=p.parent/(tmp_path.name+'.docx');d.save(target)
    result=compare_document(target,op)
    if mutation=='temperature':
        assert all(f['result']=='NOT_VERIFIABLE' for f in result['findings'] if f['field'].startswith('well.'))
    else:
        f=next(f for f in result['findings'] if f['field']=='hardware.MD (m)' and f['hardware']=='Casing')
        assert f['result']==('DISCREPANCY' if mutation=='md' else 'NOT_VERIFIABLE')
    if mutation in ('unit','duplicate'):
        assert all(f['result']=='NOT_VERIFIABLE' for f in result['findings'] if f['field'].startswith('hardware.'))


def test_directional_source_temperature_and_missing_reference_fail_closed(case):
    _,op,_=case;op=deepcopy(op)
    op['well_geometry']={'version':1,'type':'Directional','td_m':None,'survey':[
        {'MD':0,'Deviation':0,'Azimuth':0},{'MD':1000,'Deviation':30,'Azimuth':0}]}
    ref=request('well',op)['reference']['well']['full_precision']
    assert 0<ref['geo_tvd']<ref['geo_md']==1000
    op['well_geometry']['survey'][1]['MD']=0
    with pytest.raises(ValueError):request('well',op)
    op=operator();op.pop('geothermal_config')
    with pytest.raises(ValueError):request('well',op)


def test_cli_determinism_and_input_source_parent_preservation(case):
    p,op,_=case;digest=hashlib.sha256(p.read_bytes()).hexdigest()
    inputs=p.parent/'operator.json';inputs.write_text(json.dumps(op))
    outputs=[subprocess.check_output([sys.executable,'-B',str(ROOT/'c06_well.py'),str(inputs),str(p)],text=True) for _ in range(2)]
    assert outputs[0]==outputs[1] and json.loads(outputs[0])['status']=='PARTIAL_FINDINGS_ONLY'
    assert hashlib.sha256(p.read_bytes()).hexdigest()==digest
    from tools.provision import verify_candidate
    from tools.ci import source_proof
    verify_candidate(ROOT);assert source_proof()['verified_files']==143


@pytest.mark.parametrize('raw,state',[('', 'MISSING'),('0','ZERO'),('-','DASH'),('N/A','NA_PRINTED'),('bad','MALFORMED_OR_UNIT_UNPROVEN')])
def test_observed_hardware_states_are_distinct(case,monkeypatch,tmp_path,raw,state):
    p,op,source=case;monkeypatch.setattr('c06_well.request',lambda *a,**k:source)
    d=Document(p);hw=next(t for t in d.tables if t.cell(0,0).text=='Tubular and Casing Hardware')
    next(r for r in hw.rows if r.cells[0].text=='Casing').cells[3].text=raw
    target=p.parent/(tmp_path.name+'.docx');d.save(target)
    result=compare_document(target,op)
    f=next(f for f in result['findings'] if f['field']=='hardware.ID (in)' and f['hardware']=='Casing')
    assert f['reported_state']==state and f['result']!='MATCH'


@pytest.mark.parametrize('raw',['bad',0])
def test_invalid_independent_numeric_reference_cannot_match_itself(case,raw):
    p,op,_=case;op=deepcopy(op);op['hardware'][0]['ID (in)']=raw
    target=p.parent/('invalid-reference-'+str(raw)+'.docx')
    source=request('render_well',op,docx=str(target))
    assert source['reference']['geometry']['status']=='NOT_VERIFIABLE'
    result=compare_document(target,op)
    assert all(f['result']=='NOT_VERIFIABLE' for f in result['findings'] if f['field'].startswith('hardware.'))


def test_geometry_does_not_invent_an_active_cement_fluid(case):
    _,op,_=case;op=deepcopy(op)
    op['fluids']={'Spacer':{'volume':20,'density':110,'pump_rate':2}}
    ref=request('well',op)['reference']['geometry']
    assert ref['status']=='INFORMATIONAL' and ref['placement']==[]
    assert ref['client_requested_volumes_bbl']=={} and ref['client_volume_comparison']=='NOT_VERIFIABLE'


@pytest.mark.parametrize('depth,display',[(1000.05,'1000.0'),(1000.06,'1000.1')])
def test_original_scalar_md_display_boundary(case,depth,display):
    p,op,_=case;op=deepcopy(op);op['hardware'][0]['MD (m)']=depth
    target=p.parent/('md-'+str(depth)+'.docx');request('render_well',op,docx=str(target))
    f=next(f for f in compare_document(target,op)['findings'] if f['field']=='hardware.MD (m)' and f['hardware']=='Casing')
    assert f['raw_reported_text']==f['source_rounded_expected']==display
    assert f['full_precision_expected']==depth and f['result']=='MATCH'


def test_ct_zero_dimensions_remain_explicitly_unmeasured(case):
    p,_,_=case
    row=hardware('CT',500,'0',0)
    for k in ('Joint (m)','Weight (ppf)','Collapse (psi)','Burst (psi)'):row[k]=0
    op={'job_type':'CMT PLUG WITH CT','hardware':[row]}
    target=p.parent/'ct-zero.docx';request('render_well',op,docx=str(target))
    result=compare_document(target,op)
    assert next(f for f in result['findings'] if f['field']=='hardware.MD (m)')['result']=='MATCH'
    assert all(f['result']=='NOT_VERIFIABLE' for f in result['findings'] if f['field'].startswith('hardware.') and f['field'] not in ('hardware.MD (m)','hardware.Grade'))
    assert next(f for f in result['findings'] if f['field']=='hardware.ID (in)')['reported_state']=='ZERO'


def test_scalar_column_grid_must_be_proven():
    # An altered gridSpan shifts physical columns without changing printed text.
    # C01R preserves the span; consumers must not assume header indices prove units.
    from c05_fluids import compare_document
    from tests.test_c05 import operator
    with tempfile.TemporaryDirectory(dir=ROOT/'tmp',prefix='c06-grid-') as folder:
        p=Path(folder)/'synthetic.docx'
        request('render_fluids',operator(),docx=str(p))
        d=Document(p);table=next(t for t in d.tables if t.cell(0,0).text=='Fluids')
        span=OxmlElement('w:gridSpan');span.set(qn('w:val'),'2')
        table.rows[2]._tr.tc_lst[0].get_or_add_tcPr().append(span)
        p2=Path(folder)/'changed-grid.docx';d.save(p2)
        result=compare_document(p2,operator())
        assert all(f['result']=='NOT_VERIFIABLE' for f in result['findings']
                   if f.get('fluid')=='Spacer')
