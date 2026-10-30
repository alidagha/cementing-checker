"""Private synthetic correction controls, never authentic App/Golden UAT."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

import c10_corrections as c10
from test_c02 import fixture, operator


@pytest.fixture
def case():
    with tempfile.TemporaryDirectory(dir=c10.ROOT/'tmp',prefix='c10-') as folder:
        root=Path(folder); path=fixture(root/'original.docx',{'yield':'9.900','sacks':'999.9'})
        document=Document(path)
        document.sections[0].header.paragraphs[0].text='Synthetic retained header'
        document.sections[0].footer.paragraphs[0].text='Synthetic retained footer'
        document.tables[0].cell(1,1).paragraphs[0].runs[0].bold=True
        document.save(path)
        yield path, root/'copy.docx', operator()


def selection(preview,field='yield'):
    return {'schema':c10.SCHEMA,'bindings':preview['bindings'],
        'selected':[p for p in preview['proposals'] if p['field']==field]}


def findings(comparison):
    return {f['field']:f for f in comparison['findings'] if f['field'] in ('yield','sacks')}


def test_selected_copy_independent_red_green_and_unselected_discrepancy(case):
    path,out,op=case; original=path.read_bytes(); preview=c10.propose(path,op)
    old=findings(preview['comparison_before'])
    assert old['yield']['result']==old['sacks']['result']=='DISCREPANCY'
    # Independent accepted source invariant, not expected=actual from the candidate.
    assert old['yield']['source_rounded_expected']=='1.818'
    assert old['sacks']['source_rounded_expected']=='154.4'
    result=c10.apply(path,op,selection(preview),out)
    assert result['status']=='PARTIAL_FINDINGS_ONLY' and result['transaction_status']=='COPY_CREATED'
    now=findings(result['comparison_after'])
    assert now['yield']['raw_reported_text']=='1.818 cuft/sk' and now['yield']['result']=='MATCH'
    assert now['sacks']['raw_reported_text']=='999.9 sks' and now['sacks']['result']=='DISCREPANCY'
    assert result['dependency_group_complete'] is False
    assert path.read_bytes()==original and result['original_document_sha256']==hashlib.sha256(original).hexdigest()
    assert result['new_document_sha256']==hashlib.sha256(out.read_bytes()).hexdigest()
    with zipfile.ZipFile(path) as before,zipfile.ZipFile(out) as after:
        assert before.namelist()==after.namelist() and before.comment==after.comment
        changed=[n for n in before.namelist() if before.read(n)!=after.read(n)]
        assert changed==['word/document.xml']
        assert before.read('word/document.xml').replace(b'9.900 cuft/sk',b'1.818 cuft/sk',1)==after.read('word/document.xml')
        for a,b in zip(before.infolist(),after.infolist()):
            assert (a.filename,a.date_time,a.extra,a.comment,a.external_attr,a.flag_bits)==(b.filename,b.date_time,b.extra,b.comment,b.external_attr,b.flag_bits)
    assert not list(out.parent.glob('.c10-*'))


def test_multiple_explicit_cells_and_repeat_determinism(case):
    path,out,op=case; p=c10.propose(path,op)
    assert c10.propose(path,op)==p
    s={'schema':c10.SCHEMA,'bindings':p['bindings'],'selected':p['proposals']}
    a=c10.apply(path,op,s,out); b=c10.apply(path,op,s,out.with_name('second.docx'))
    assert a['transaction_status']==b['transaction_status']=='COPY_CREATED'
    assert out.read_bytes()==out.with_name('second.docx').read_bytes()
    assert all(f['result']=='MATCH' for f in findings(a['comparison_after']).values())
    assert a['status']=='PARTIAL_FINDINGS_ONLY'


@pytest.mark.parametrize('mutation',['document','operator','source','template','finding','locator','slurry','replacement','empty','duplicate'])
def test_stale_forged_or_implicit_selection_never_writes(case,mutation):
    path,out,op=case; p=c10.propose(path,op); s=selection(p)
    if mutation=='document':
        d=Document(path); d.add_paragraph('Untrusted subsequent change');d.save(path)
    elif mutation=='operator':op['slurries']['Main']['client_volume_bbl']=60
    elif mutation=='source':s['bindings']['source_sha']='0'*40
    elif mutation=='template':s['bindings']['template_sha256']='0'*64
    elif mutation=='finding':s['selected'][0]['finding_id']='0'*64
    elif mutation=='locator':s['selected'][0]['locator']['row']=2
    elif mutation=='slurry':s['selected'][0]['slurry']='Tail'
    elif mutation=='replacement':s['selected'][0]['replacement']='1.818 bbl'
    elif mutation=='empty':s['selected']=[]
    elif mutation=='duplicate':s['selected']*=2
    original=path.read_bytes(); result=c10.apply(path,op,s,out)
    assert result['status']=='INPUT_REJECTED' and result['transaction_status']=='NO_COPY_CREATED'
    assert not out.exists() and not list(out.parent.glob('.c10-*')) and path.read_bytes()==original


@pytest.mark.parametrize('mutation',['split','multi-paragraph','field','hidden','numbering','inherited-numbering',
    'hyperlink','control','merged','duplicate','excluded','unknown-style'])
def test_unsafe_physical_targets_require_manual_correction(case,mutation):
    path,out,op=case; d=Document(path); cell=d.tables[0].cell(1,1); p=cell.paragraphs[0]
    if mutation=='split':
        p.runs[0].text='9.';p.add_run('900 cuft/sk').bold=True
    elif mutation=='multi-paragraph':cell.add_paragraph('Unproved extra claim')
    elif mutation=='field':p._p.append(OxmlElement('w:fldSimple'))
    elif mutation=='hidden':p.runs[0]._r.get_or_add_rPr().append(OxmlElement('w:vanish'))
    elif mutation=='numbering':p._p.get_or_add_pPr().append(OxmlElement('w:numPr'))
    elif mutation=='inherited-numbering':
        from docx.enum.style import WD_STYLE_TYPE
        style=d.styles.add_style('C10Numbered',WD_STYLE_TYPE.PARAGRAPH)
        style.element.get_or_add_pPr().append(OxmlElement('w:numPr'));p.style=style
    elif mutation=='hyperlink':
        link=OxmlElement('w:hyperlink');p._p.append(link)
    elif mutation=='control':p._p.append(OxmlElement('w:sdt'))
    elif mutation=='merged':cell._tc.get_or_add_tcPr().append(OxmlElement('w:vMerge'))
    elif mutation=='duplicate':
        row=deepcopy(d.tables[0].rows[1]._tr);d.tables[0]._tbl.append(row)
    elif mutation=='excluded':d.paragraphs[2].text='Document Control'
    elif mutation=='unknown-style':
        style=OxmlElement('w:rStyle');style.set(qn('w:val'),'UnresolvedC10Style');p.runs[0]._r.get_or_add_rPr().append(style)
    d.save(path); original=path.read_bytes(); preview=c10.propose(path,op)
    if preview['status']=='PARTIAL_FINDINGS_ONLY':
        offers=[x for x in preview['proposals'] if x['field']=='yield']
        assert not any(x['action']=='SELECTABLE' for x in offers)
        result=c10.apply(path,op,selection(preview),out)
        assert result['status']=='INPUT_REJECTED'
    else:assert preview['status']=='INPUT_REJECTED'
    assert path.read_bytes()==original and not out.exists()


@pytest.mark.parametrize('destination',['original','existing','symlink','escape','symlink-parent'])
def test_output_path_safety_and_no_partial_files(case,destination):
    path,out,op=case;p=c10.propose(path,op)
    if destination=='original':out=path
    elif destination=='existing':out.write_bytes(b'keep existing')
    elif destination=='symlink':out.symlink_to(path)
    elif destination=='escape':out=c10.ROOT/'outside-output.docx'
    elif destination=='symlink-parent':
        alias=path.parent/'alias';alias.symlink_to(path.parent,target_is_directory=True);out=alias/'copy.docx'
    initial=out.read_bytes() if out.exists() else None
    result=c10.apply(path,op,selection(p),out)
    assert result['status']=='INPUT_REJECTED'
    assert (out.read_bytes() if out.exists() else None)==initial and not list(path.parent.glob('.c10-*'))


def test_damaged_docx_and_original_symlink_reject(case):
    path,out,op=case; alias=path.parent/'alias.docx';alias.symlink_to(path)
    assert c10.propose(alias,op)['status']=='INPUT_REJECTED'
    path.write_bytes(b'not a Word package')
    assert c10.propose(path,op)['status']=='INPUT_REJECTED'
    assert not out.exists()


def test_recheck_failure_cleans_staged_copy(case,monkeypatch):
    path,out,op=case
    preview=c10.propose(path,op)
    def unavailable(*args):raise c10.Refusal('RECHECK_UNAVAILABLE_NOT_VERIFIABLE')
    monkeypatch.setattr(c10,'recheck',unavailable)
    with pytest.raises(c10.Refusal,match='RECHECK_UNAVAILABLE'):
        c10.execute(path,op,selection(preview),out)
    assert not out.exists() and not list(path.parent.glob('.c10-*'))


def test_unpreservable_zip_compression_metadata_refuses_copy(case):
    path,out,op=case
    from copy import copy
    with zipfile.ZipFile(path) as z:
        members=[(copy(i),z.read(i)) for i in z.infolist()]
    with zipfile.ZipFile(path,'w') as z:
        for info,raw in members:z.writestr(info,raw,compresslevel=0)
    original=path.read_bytes();p=c10.propose(path,op)
    result=c10.apply(path,op,selection(p),out)
    assert result['status']=='INPUT_REJECTED' and not out.exists()
    assert result['reason']=='UNSELECTED_ZIP_METADATA_CANNOT_BE_PRESERVED'
    assert path.read_bytes()==original and not list(path.parent.glob('.c10-*'))


def test_cli_exact_selection_json_binding_and_source_rendered_template(case):
    path,out,op=case
    from c03.c03_source import request,HERE
    with tempfile.TemporaryDirectory(dir=HERE/'tmp',prefix='c10-source-') as folder:
        rendered=Path(folder)/'synthetic.docx'
        source=request({'action':'render_fixture','operator':op,'docx':str(rendered),
            'fixture':{'Main':{'dead':0,'tank':'Mud Reserve Tanks'}}})
        assert source['source_sha']==c10.SOURCE_SHA
        d=Document(rendered)
        for table in d.tables:
            if table.rows and table.cell(0,0).text=='Main Cement Slurry Data':
                for row in table.rows:
                    for i,cell in enumerate(row.cells[:-1]):
                        if cell.text=='Yield':row.cells[i+1].text='9.900 cuft/sk'
        d.save(path)
    original=path.read_bytes();opfile=path.parent/'operator.json';opfile.write_text(json.dumps(op))
    cmd=[sys.executable,'-B',str(c10.ROOT/'c10_corrections.py'),str(opfile),str(path)]
    p=subprocess.run(cmd,capture_output=True,text=True);assert p.returncode==0,p.stderr
    preview=json.loads(p.stdout);s=selection(preview)
    assert len(s['selected'])==1 and s['selected'][0]['action']=='SELECTABLE'
    sf=path.parent/'selection.json';sf.write_text(json.dumps(s))
    run=subprocess.run(cmd+['--selection',str(sf),'--output',str(out)],capture_output=True,text=True)
    assert run.returncode==0,run.stdout+run.stderr
    result=json.loads(run.stdout);assert result['transaction_status']=='COPY_CREATED'
    assert result['comparison_after']['parser_schema']=='c01.1' and path.read_bytes()==original


def test_runtime_write_network_and_source_boundaries(case):
    path,out,op=case
    script='''
import socket
from engineering_source import install_audit
from comparison_records import ROOT
install_audit()
denied=0
for f in (lambda:socket.getaddrinfo('example.com',443),lambda:(ROOT/'forbidden.txt').write_text('no')):
    try:f()
    except RuntimeError:denied+=1
assert denied==2
'''
    p=subprocess.run([sys.executable,'-B','-c',script],cwd=c10.ROOT,capture_output=True,text=True)
    assert p.returncode==0,p.stderr
    assert not (c10.ROOT/'forbidden.txt').exists()
    preview=c10.propose(path,op)
    assert preview['bindings']['template_sha256']==c10.TEMPLATE_SHA
    assert preview['bindings']['source_sha']==c10.SOURCE_SHA
    assert len(preview['comparison_before']['reference']['source_module_sha256'])>5


def test_two_slurries_select_lead_keeps_tail_and_notes(case):
    path,out,op=case
    from c03.c03_source import request,HERE
    recipe=op['slurries'].pop('Main');op['active_fluids']=['Lead','Tail']
    op['slurries']={'Lead':deepcopy(recipe),'Tail':deepcopy(recipe)}
    op['slurries']['Tail']['density_pcf']=118
    with tempfile.TemporaryDirectory(dir=HERE/'tmp',prefix='c10-two-') as folder:
        rendered=Path(folder)/'synthetic.docx'
        request({'action':'render_fixture','operator':op,'docx':str(rendered),
            'fixture':{'Lead':{'dead':0,'tank':'Mud Reserve Tanks'},'Tail':{'dead':0,'tank':'Cement Unit Tanks'}}})
        d=Document(rendered)
        for table in d.tables:
            if table.rows and table.cell(0,0).text in ('Lead Cement Slurry Data','Tail Cement Slurry Data'):
                for row in table.rows:
                    for i,cell in enumerate(row.cells[:-1]):
                        if cell.text=='Yield':row.cells[i+1].text='9.900 cuft/sk'
        d.save(path)
    before=path.read_bytes();p=c10.propose(path,op)
    offers=[f for f in p['proposals'] if f['field']=='yield']
    assert len(offers)==2 and {f['replacement'] for f in offers}=={'1.818 cuft/sk','1.360 cuft/sk'}
    s={'schema':c10.SCHEMA,'bindings':p['bindings'],'selected':[f for f in offers if f['slurry']=='Lead']}
    r=c10.apply(path,op,s,out);assert r['transaction_status']=='COPY_CREATED'
    found={f['slurry']:f for f in r['comparison_after']['findings'] if f['field']=='yield'}
    assert found['Lead']['result']=='MATCH' and found['Tail']['result']=='DISCREPANCY'
    assert found['Tail']['raw_reported_text']=='9.900 cuft/sk' and path.read_bytes()==before


@pytest.mark.parametrize('job,fluid',[('CMT PLUG WITH CT','Main'),('CMT SQUEEZE','Main'),('CSG 20"','Scavenger')])
def test_special_roles_are_not_selectable(case,job,fluid):
    path,out,op=case;op['job_type']=job
    if fluid!='Main':
        op['active_fluids']=[fluid];op['slurries']={fluid:op['slurries']['Main']};fixture(path,{'yield':'9.900'},fluid=fluid)
    p=c10.propose(path,op)
    assert p['status']=='INPUT_REJECTED' or not any(x['action']=='SELECTABLE' for x in p['proposals'])
    assert not out.exists()
