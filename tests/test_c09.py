"""C09 genuine source projection and independent procedure/NOTE mutations."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile

from docx import Document
import pytest

from checker_parser import parse_docx
from comparison_records import ROOT
from narrative_source import request
from c09_procedure import compare_document
from test_c08 import operator as program


def operator():
    p=program();p['cement_operator']['slurries']['Main']['additives']=[
        {'Material Type':'Anti Foam','Name':'TA-47','Physical State':'Liquid','Mix Method':'In Mix Water','User Input':0.03},
        {'Material Type':'Custom Powder','Name':'Synthetic Powder','Physical State':'Powder','Mix Method':'In Mix Water','User Input':1,'Density':1.4}]
    return {'program':p,'preparation':{'Main':{'tank_name':'Mud Reserve Tanks','dead_volume_bbl':5,'surface_hardened_hours':8}},
        'well_geometry':{'version':1,'type':'Vertical','td_m':1000}}


@pytest.fixture(scope='module')
def case():
    with tempfile.TemporaryDirectory(dir=ROOT/'tmp',prefix='c09-') as folder:
        p=Path(folder)/'synthetic.docx';op=operator();source=request(op,'render_procedure',str(p))
        yield p,op,source


def test_actual_source_state_signature_and_independent_wording_oracles(case):
    p,op,source=case;ref=source['reference'];text=ref['official_generated_text']
    assert ref['approval_flags_supplied'] is False and ref['session_state_implementation']=='streamlit.runtime.state.session_state_proxy'
    assert ref['procedure_source_signature']
    assert {'phase_10_procedure.py:generate_official_procedure','phase_10_procedure.py:build_ordered_notes',
        'phase_10_procedure.py:synchronize_report_texts','phase_5_cement.py:refresh_cement_calculations'}<=set(ref['source_calls'])
    assert '1. Clean up Mud Reserve Tanks and preparing chemicals for mixing fluid.' in text
    assert 'Mix and pump 50.0 bbl Main cement slurry (104 pcf).' in text
    assert 'Add TA-47 to Fresh Water and allow mixing for 10 min.' in text
    assert 'Add Synthetic Powder and allow mixing until the cement operation will finish.' in text
    assert 'NOTE 1: The Calculated Temperature is based on True Vertical Depth.' in ref['all_source_notes']
    assert any('Surface Sample: Hardened Condition Observed After 8 Hours' in n for n in ref['all_source_notes'])
    result=compare_document(p,op)
    assert result['status']=='PARTIAL_FINDINGS_ONLY' and result['findings']
    assert any(f['result']=='MATCH' and f['field']=='procedure' for f in result['findings'])
    assert all(f['result']!='DISCREPANCY' for f in result['findings'])
    assert any(f['field']=='procedure' and f['result']=='NOT_VERIFIABLE' and 'rig-site' in f['reason'] for f in result['findings'])
    assert all(f['result']=='MATCH' for f in result['findings'] if f['field']=='ordered_notes')
    assert any(f['slurry']=='Main' and 'Mix above' in str(f['source_wording']) for f in result['findings'])
    assert result['records']==parse_docx(p)['records'] and result['unverified_record_orders']


@pytest.mark.parametrize('mutation',['wording','numbering','reorder','omit','duplicate','tank','water','mix-time','cross-fluid','cross-section','unknown-heading'])
def test_untrusted_procedure_mutations_never_false_match(case,monkeypatch,mutation):
    p,op,source=case;monkeypatch.setattr('c09_procedure.request',lambda *a:source)
    d=Document(p);start=next(i for i,x in enumerate(d.paragraphs) if x.text.strip()=='X. Cementing Procedure' or x.text.strip()=='IX. Cementing Procedure')
    paragraphs=[x for x in d.paragraphs[start+1:] if x.text.strip()]
    first=next(x for x in paragraphs if x.text.startswith('1. '))
    if mutation=='wording':first.text=first.text.replace('Clean up','Do not clean')
    elif mutation=='numbering':first.text=first.text.replace('1. ','99. ',1)
    elif mutation=='reorder':first._p.addprevious(deepcopy(paragraphs[1]._p));paragraphs[1]._p.getparent().remove(paragraphs[1]._p)
    elif mutation=='omit':first._p.getparent().remove(first._p)
    elif mutation=='duplicate':first._p.addnext(deepcopy(first._p))
    elif mutation in ('tank','water'):
        para=next(x for x in paragraphs if 'Fill up' in x.text)
        para.text=para.text.replace('Mud Reserve Tanks','Cement Unit Tanks') if mutation=='tank' else para.text.replace(' bbl','0 bbl',1)
    elif mutation=='mix-time':
        para=next(x for x in paragraphs if '10 min.' in x.text);para.text=para.text.replace('10 min.','100 min.')
    elif mutation=='cross-fluid':
        para=next(x for x in paragraphs if 'Mix and pump' in x.text);para.text=para.text.replace('Main','Tail')
    elif mutation=='cross-section':d.paragraphs[start].text='Executive Summary'
    else:first._p.get_or_add_pPr().get_or_add_pStyle().val='Heading2';first.text='Unproved subsection'
    target=p.parent/(mutation+'.docx');d.save(target)
    r=compare_document(target,op)
    assert any(f['result'] in ('DISCREPANCY','NOT_VERIFIABLE') for f in r['findings'])
    assert not any(f['field']=='procedure.sequence' and f['result']=='MATCH' for f in r['findings'])


@pytest.mark.parametrize('mutation',['duplicate','omit','renumber','tank','water','hours','cross-section','break-group'])
def test_note_identity_order_and_source_group_negatives(case,monkeypatch,mutation):
    p,op,source=case;monkeypatch.setattr('c09_procedure.request',lambda *a:source)
    d=Document(p);para=next(p for p in d.paragraphs if 'Mix above Additives in' in p.text)
    if mutation=='duplicate':para._p.addnext(deepcopy(para._p))
    elif mutation=='omit':para._p.getparent().remove(para._p)
    elif mutation=='renumber':para.text=para.text.replace('NOTE 5:','NOTE 55:')
    elif mutation=='tank':para.text=para.text.replace('Mud Reserve Tanks','Cement Unit Tanks')
    elif mutation=='water':para.text=para.text.replace(' bbl','0 bbl',1)
    elif mutation=='hours':
        para=next(p for p in d.paragraphs if 'Surface Sample:' in p.text);para.text=para.text.replace('8 Hours','80 Hours')
    elif mutation=='cross-section':
        proc=next(p for p in d.paragraphs if p.text.endswith('Cementing Procedure'));proc._p.addnext(deepcopy(para._p));para._p.getparent().remove(para._p)
    else:
        para.insert_paragraph_before('Unrelated unproved paragraph')
    target=p.parent/('note-'+mutation+'.docx');d.save(target);r=compare_document(target,op)
    assert not any(f['field']=='ordered_notes.sequence' and f['result']=='MATCH' for f in r['findings'])
    assert any(f['result'] in ('DISCREPANCY','NOT_VERIFIABLE') for f in r['findings'])


@pytest.mark.parametrize('bad',['missing','duplicate-role','negative','nan','tank','lab','ct','squeeze-incomplete','preflush','spacer'])
def test_incomplete_independent_context_and_unsupported_paths(case,bad):
    p,op,_=case;op=deepcopy(op)
    if bad=='missing':del op['preparation']
    elif bad=='duplicate-role':op['preparation']['Tail']=op['preparation']['Main']
    elif bad=='negative':op['preparation']['Main']['dead_volume_bbl']=-1
    elif bad=='nan':op['preparation']['Main']['dead_volume_bbl']=float('nan')
    elif bad=='tank':op['preparation']['Main']['tank_name']='Invented Tank'
    elif bad=='lab':del op['preparation']['Main']['surface_hardened_hours']
    elif bad=='ct':op['program']['job_type']='CMT PLUG WITH CT'
    elif bad=='squeeze-incomplete':op['program']['job_type']='CMT SQUEEZE'
    else:
        fluid='Pre Flush' if bad=='preflush' else 'Spacer'
        op['program']['fluids'][fluid]={'volume':10,'density':65,'pump_rate':2}
    result=compare_document(p,op)
    assert all(f['result']=='NOT_VERIFIABLE' for f in result['findings'])
    assert len(result['unverified_record_orders'])==len(result['records'])


def test_two_slurries_distinct_tanks_and_operational_spacer_source_projection():
    op=operator();p=op['program'];p['job_type']='CSG 7"';job=p['job_type']
    p['placement_config']['job_type']=job;p['executive_summary_config']['job_type']=job
    original=p['cement_operator']['slurries']['Main']
    p['fluids']={'Lead':{'volume':40,'density':104,'pump_rate':2},
        'Tail':{'volume':10,'density':118,'pump_rate':2},
        'Spacer Behind':{'volume':5,'density':70,'pump_rate':2,'material_name':'Synthetic operational spacer'}}
    p['cement_operator']={'job_type':job,'active_fluids':['Lead','Tail','Spacer Behind'],
        'slurries':{'Lead':{**original,'client_volume_bbl':40},'Tail':{**original,'client_volume_bbl':10,'density_pcf':118}}}
    p['tops']={n:{'top_job_type':job,'top_mode':'Depth (m MD)','top_depth':900} for n in ('Lead','Tail')}
    p['executive_summary_config']['per_slurry']={n:{'volume_basis':'client_request'} for n in ('Lead','Tail')}
    op['preparation']={'Lead':{'tank_name':'Mud Reserve Tanks','dead_volume_bbl':5,'surface_hardened_hours':8},
        'Tail':{'tank_name':'Batch Mixer Tanks','dead_volume_bbl':2,'surface_hardened_hours':10}}
    with tempfile.TemporaryDirectory(dir=ROOT/'tmp') as folder:
        path=Path(folder)/'two-slurry.docx';source=request(op,'render_procedure',str(path));r=compare_document(path,op)
        assert not any(f['result']=='DISCREPANCY' for f in r['findings'])
        notes=[f for f in r['findings'] if f['field']=='ordered_notes' and 'Mix above' in str(f['source_wording'])]
        assert {f['slurry'] for f in notes}=={'Lead','Tail'} and all(f['result']=='MATCH' for f in notes)
        assert 'Spacer Behind' in source['reference']['official_generated_text']
        # The aggregate preparation instruction is verified as exact source text,
        # never apportioned/inverted to infer individual allowance/tank proof.
        assert 'Lead:' in source['reference']['official_generated_text'] and 'Tail:' in source['reference']['official_generated_text']


def test_manual_procedure_independent_approval_and_no_source_gate_override(case):
    p,op,_=case;op=deepcopy(op);text='1. Independent reviewed synthetic instruction.\n2. Confirm the independently reviewed plan.'
    op['manual_procedure']={'text':text,'approval_reference':'SYNTHETIC independent approval record','text_sha256':hashlib.sha256(text.encode()).hexdigest()}
    with tempfile.TemporaryDirectory(dir=ROOT/'tmp') as folder:
        target=Path(folder)/'manual.docx';source=request(op,'render_procedure',str(target));r=compare_document(target,op)
        assert all(f['result']=='MATCH' for f in r['findings'] if f['field']=='procedure')
        assert source['reference']['kind']=='INDEPENDENT_APPROVED_MANUAL_TEXT' and source['reference']['approval_flags_supplied'] is False
        op['manual_procedure']['approval_reference']=''
        assert all(f['result']=='NOT_VERIFIABLE' for f in compare_document(target,op)['findings'])


def test_cli_determinism_preserves_document_operator_and_raw_records(case):
    p,op,_=case;operator_path=p.parent/'operator.json';operator_path.write_text(json.dumps(op));before=(p.read_bytes(),operator_path.read_bytes())
    cmd=[sys.executable,'-B',str(ROOT/'c09_procedure.py'),str(p),str(operator_path)]
    a=subprocess.run(cmd,text=True,capture_output=True,check=True);b=subprocess.run(cmd,text=True,capture_output=True,check=True)
    assert a.stdout==b.stdout
    result=json.loads(a.stdout);assert result['status']=='PARTIAL_FINDINGS_ONLY' and result['records']==parse_docx(p)['records']
    assert before==(p.read_bytes(),operator_path.read_bytes())


@pytest.mark.parametrize('location',['procedure','note','summary'])
@pytest.mark.parametrize('mechanism',['numbering','field','hidden'])
def test_automatic_word_numbering_fails_closed(case,monkeypatch,location,mechanism):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from c08_summary import compare_document as compare_summary
    p,op,source=case;monkeypatch.setattr('c09_procedure.request',lambda *a:source)
    d=Document(p)
    if location=='procedure':para=next(p for p in d.paragraphs if p.text.startswith('1. '))
    elif location=='note':para=next(p for p in d.paragraphs if 'Mix above Additives in' in p.text)
    else:para=next(p for p in d.paragraphs if p.text.startswith('Enclosed'))
    if mechanism=='numbering':
        numbering=OxmlElement('w:numPr');identifier=OxmlElement('w:numId');identifier.set(qn('w:val'),'987654')
        numbering.append(identifier);para._p.get_or_add_pPr().append(numbering)
    elif mechanism=='field':para._p.append(OxmlElement('w:fldSimple'))
    else:para.runs[0]._r.get_or_add_rPr().append(OxmlElement('w:vanish'))
    target=p.parent/(mechanism+'-'+location+'.docx');d.save(target)
    r=compare_summary(target,op['program']) if location=='summary' else compare_document(target,op)
    assert all(f['result']=='NOT_VERIFIABLE' for f in r['findings'])
