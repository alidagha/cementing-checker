"""Independent synthetic reference changes; no Word repair or Golden claim."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile

from docx import Document
from docx.oxml import OxmlElement
import pytest

import g06a_client_volume as g
from test_c02 import operator, fixture


@pytest.fixture
def case():
    with tempfile.TemporaryDirectory(dir=g.ROOT/'tmp',prefix='g06a-') as folder:
        path=fixture(Path(folder)/'original.docx')
        yield path,operator()


def request(path,op,volume=55,slurry='Main'):
    inspection=g.inspect_reference(path,op,slurry)
    assert inspection['status']=='PARTIAL_FINDINGS_ONLY',inspection
    return {'schema':g.SCHEMA,'bindings':inspection['bindings'],'slurry':slurry,'client_volume_bbl':volume}


def dependencies(result):return {x['dependency']:x for x in result['source_dependency_ledger']}


def test_main_volume_changes_expectation_not_word_with_independent_oracles(case):
    path,op=case; original=path.read_bytes();old_op=deepcopy(op);edit=request(path,op)
    result=g.volume_impact(path,op,edit)
    assert result['status']=='PARTIAL_FINDINGS_ONLY' and result['operation']=='CLIENT_VOLUME_REFERENCE_IMPACT'
    rows=dependencies(result)
    assert rows['sacks']['old_source_display']=='154.4' and rows['sacks']['new_source_display']=='169.9'
    assert rows['pure_mix_water']['old_source_display']=='34.8' and rows['pure_mix_water']['new_source_display']=='38.3'
    assert rows['yield']['old_source_display']==rows['yield']['new_source_display']=='1.818'
    assert rows['calculated_600ml_recipe']['classification']=='SOURCE_UNCHANGED'
    # Accepted 104-pcf/SG3.2 independent invariant; not extracted from Word.
    assert rows['calculated_600ml_recipe']['new_source_display'][0]['Mass']=='582.0'
    sacks=next(x for x in result['word_claim_impacts'] if x['field']=='sacks')
    assert sacks['old_comparison']=='MATCH' and sacks['new_comparison']=='DISCREPANCY'
    assert sacks['original_raw_claim']=='154.4 sks' and sacks['classification']=='SOURCE_CHANGED'
    assert result['original_comparison']['findings']!=result['edited_comparison']['findings']
    assert op==old_op and result['edited_operator_reference']['slurries']['Main']['client_volume_bbl']==55
    assert result['bindings']['original_operator_sha256']==g.digest(old_op)
    assert result['edited_operator_sha256']==g.digest(result['edited_operator_reference'])
    assert result['original_document_sha256']==result['unchanged_document_sha256']==hashlib.sha256(original).hexdigest()
    assert path.read_bytes()==original and result['word_modified'] is result['docx_written'] is False
    assert list(path.parent.iterdir())==[path]


def test_reference_correction_can_remove_discrepancy_without_claiming_docx_repair(case):
    path,op=case;fixture(path,{'sacks':'169.9','yield':'9.900'});original=path.read_bytes()
    result=g.volume_impact(path,op,request(path,op))
    mapped={x['field']:x for x in result['word_claim_impacts'] if x['field'] in ('sacks','yield')}
    assert mapped['sacks']['old_comparison']=='DISCREPANCY' and mapped['sacks']['new_comparison']=='MATCH'
    assert mapped['yield']['old_comparison']==mapped['yield']['new_comparison']=='DISCREPANCY'
    assert result['status']=='PARTIAL_FINDINGS_ONLY' and result['word_modified'] is False
    assert path.read_bytes()==original


def test_noop_keeps_source_and_unresolved_dependencies(case):
    path,op=case;r=g.volume_impact(path,op,request(path,op,50))
    assert r['status']=='PARTIAL_FINDINGS_ONLY'
    assert all(x['classification']=='SOURCE_UNCHANGED' for x in r['source_dependency_ledger'])
    assert r['bindings']['original_operator_sha256']==r['edited_operator_sha256']
    assert {x['classification'] for x in r['unresolved_dependencies']}=={'NOT_VERIFIABLE','MANUAL_REVIEW_REQUIRED'}
    assert 'total_preparation_water_and_gross_with_dead_volume' not in dependencies(r)
    assert r['invariants']['measured_results_refreshed'] is False


def test_source_precision_change_can_be_invisible_at_original_display_precision(case):
    path,op=case;r=g.volume_impact(path,op,request(path,op,50.00001));rows=dependencies(r)
    sacks=rows['sacks']
    assert sacks['classification']=='SOURCE_CHANGED' and sacks['old_full_precision']!=sacks['new_full_precision']
    assert sacks['old_source_display']==sacks['new_source_display']=='154.4' and sacks['source_display_changed'] is False
    assert rows['client_requested_volume']['old_source_display'] is None and rows['client_requested_volume']['source_display_changed'] is None
    claim=next(x for x in r['word_claim_impacts'] if x['field']=='sacks')
    assert claim['old_comparison']==claim['new_comparison']=='MATCH' and r['word_modified'] is False
    assert r['status']=='PARTIAL_FINDINGS_ONLY'


@pytest.mark.parametrize('value',[0,-1,True,'55',None,float('nan'),float('inf')])
def test_invalid_volume_never_changes_inputs(case,value):
    path,op=case;original=path.read_bytes();edit=request(path,op);edit['client_volume_bbl']=value
    if isinstance(value,float) and value!=value or value==float('inf'):
        with pytest.raises(ValueError):g.volume_impact(path,op,edit)
    else:assert g.volume_impact(path,op,edit)['status']=='INPUT_REJECTED'
    assert path.read_bytes()==original and op['slurries']['Main']['client_volume_bbl']==50


@pytest.mark.parametrize('mutation',['document','operator','reference-hash','source-pin','template-pin','swapped-slurry','missing-selection','duplicate-selection','measurement'])
def test_stale_forged_or_unsupported_request_rejected(case,mutation):
    path,op=case;edit=request(path,op)
    if mutation=='document':
        d=Document(path);d.add_paragraph('Synthetic subsequent change');d.save(path)
    elif mutation=='operator':op['slurries']['Main']['density_pcf']=105
    elif mutation=='reference-hash':edit['bindings']['original_reference_sha256']='0'*64
    elif mutation=='source-pin':edit['bindings']['source_sha']='0'*40
    elif mutation=='template-pin':edit['bindings']['template_sha256']='0'*64
    elif mutation=='swapped-slurry':edit['slurry']='Tail'
    elif mutation=='missing-selection':del edit['slurry']
    elif mutation=='duplicate-selection':edit['slurry']=['Main','Main']
    elif mutation=='measurement':edit['measurements']={'Main':{'free_water':0}}
    before=path.read_bytes();r=g.volume_impact(path,op,edit)
    assert r['status']=='INPUT_REJECTED' and not r['word_modified'] and path.read_bytes()==before


@pytest.mark.parametrize('mutation',['duplicate','hidden','split','corrupt','wrong-role'])
def test_unproved_physical_identity_fails_closed(case,mutation):
    path,op=case
    if mutation=='duplicate':fixture(path,duplicate=True)
    elif mutation=='corrupt':path.write_bytes(b'not a DOCX')
    elif mutation=='wrong-role':fixture(path,fluid='Tail')
    else:
        d=Document(path);p=d.tables[0].cell(1,1).paragraphs[0]
        if mutation=='hidden':p.runs[0]._r.get_or_add_rPr().append(OxmlElement('w:vanish'))
        else:p.runs[0].text='1.';p.add_run('818 cuft/sk')
        d.save(path)
    before=path.read_bytes();r=g.inspect_reference(path,op,'Main')
    assert r['status']=='INPUT_REJECTED' and path.read_bytes()==before


@pytest.mark.parametrize('job,slurry',[('CMT PLUG WITH CT','Main'),('CMT SQUEEZE','Main'),('CSG 20"','Scavenger'),('UNKNOWN','Main')])
def test_unsupported_job_or_selected_role(case,job,slurry):
    path,op=case;op['job_type']=job
    if slurry!='Main':op['active_fluids']=[slurry];op['slurries']={slurry:op['slurries']['Main']}
    assert g.inspect_reference(path,op,slurry)['status']=='INPUT_REJECTED'


def test_two_slurries_unselected_source_and_word_claims_are_identical(case):
    path,op=case;recipe=op['slurries'].pop('Main');op['active_fluids']=['Lead','Tail']
    op['slurries']={'Lead':deepcopy(recipe),'Tail':dict(recipe,client_volume_bbl=30,density_pcf=118)}
    fixture(path,fluid='Lead');second=fixture(path.parent/'tail-control.docx',fluid='Tail')
    d=Document(path);other=Document(second);note=next(p for p in d.paragraphs if p.text.startswith('NOTE'))
    for table in other.tables[:2]:note._p.addprevious(deepcopy(table._tbl))
    d.save(path);before=path.read_bytes();edit=request(path,op,55,'Lead')
    r=g.volume_impact(path,op,edit);assert r['status']=='PARTIAL_FINDINGS_ONLY'
    old=[f for f in r['original_comparison']['findings'] if f['slurry']=='Tail']
    new=[f for f in r['edited_comparison']['findings'] if f['slurry']=='Tail']
    assert old and old==new and r['invariants']['other_slurries'] is True
    assert r['edited_operator_reference']['slurries']['Tail']==op['slurries']['Tail']
    assert path.read_bytes()==before
    crossed=deepcopy(edit);crossed['slurry']='Tail'
    assert g.volume_impact(path,op,crossed)['status']=='INPUT_REJECTED'


def test_source_additives_salt_and_solution_remain_distinct_from_actual_dv(case):
    path,op=case
    op['slurries']['Main']['additives']=[{'Material Type':'Custom Powder','Name':'Synthetic powder',
        'Physical State':'Powder','Mix Method':'In Mix Water','User Input':1,'Density':1.4},
        {'Material Type':'Custom Liquid','Name':'Synthetic liquid','Physical State':'Liquid',
         'Mix Method':'In Mix Water','User Input':0.1,'Density':1.1},
        {'Material Type':'Custom Powder','Name':'Synthetic dry blend','Physical State':'Powder',
         'Mix Method':'Dry Blend','User Input':2,'Density':2.5},
        {'Material Type':'Nacl','Name':'SALT','Physical State':'Powder',
         'Mix Method':'In Mix Water','User Input':5,'Density':2.16}]
    r=g.volume_impact(path,op,request(path,op));rows=dependencies(r)
    assert rows['base_field_materials']['classification']=='SOURCE_CHANGED'
    assert rows['calculated_600ml_recipe']['classification']=='SOURCE_UNCHANGED'
    assert rows['base_field_materials']['new_full_precision']['salt_field_lb']>rows['base_field_materials']['old_full_precision']['salt_field_lb']>0
    assert rows['solution_not_pure_water']['new_full_precision']!=rows['pure_mix_water']['new_full_precision']
    assert all('lbs or gal (with dead Vol.)' not in row for row in rows['base_field_materials']['new_source_display']['additives'])
    assert 'total_preparation_water_and_gross_with_dead_volume' not in rows
    assert 'no actual DV inferred or carried' in r['reference_dead_volume_policy']


def test_private_paths_cli_determinism_and_duplicate_json(case):
    path,op=case;opfile=path.parent/'operator.json';opfile.write_text(json.dumps(op))
    cmd=[sys.executable,'-B',str(g.ROOT/'g06a_client_volume.py'),str(opfile),str(path),'--slurry','Main']
    a=subprocess.run(cmd,capture_output=True,text=True);b=subprocess.run(cmd,capture_output=True,text=True)
    assert a.returncode==b.returncode==0 and json.loads(a.stdout)==json.loads(b.stdout)
    inspection=json.loads(a.stdout);edit={'schema':g.SCHEMA,'bindings':inspection['bindings'],'slurry':'Main','client_volume_bbl':55}
    ef=path.parent/'edit.json';ef.write_text(json.dumps(edit))
    completed=subprocess.run(cmd[:-2]+['--edit',str(ef)],capture_output=True,text=True)
    assert completed.returncode==0 and json.loads(completed.stdout)['operation']=='CLIENT_VOLUME_REFERENCE_IMPACT'
    ef.write_text('{"slurry":"Main","slurry":"Tail"}')
    bad=subprocess.run(cmd[:-2]+['--edit',str(ef)],capture_output=True,text=True)
    assert bad.returncode==2 and json.loads(bad.stdout)['status']=='INPUT_REJECTED'
    alias=path.parent/'alias.docx';alias.symlink_to(path)
    assert g.inspect_reference(alias,op,'Main')['status']=='INPUT_REJECTED'
    assert g.inspect_reference(g.ROOT/'snapshot/master_template.docx',op,'Main')['status']=='INPUT_REJECTED'


def test_original_states_unmapped_claims_and_missing_measurements_not_coerced(case):
    path,op=case;fixture(path,{'sacks':'0','yield':'1.818'})
    r=g.volume_impact(path,op,request(path,op));before=r['original_comparison']['findings'];after=r['edited_comparison']['findings']
    assert len(before)==len(after) and any(x['result']=='UNIMPLEMENTED' for x in before)
    for a,b in zip(before,after):assert a['raw_reported_text']==b['raw_reported_text'] and a['reported_state']==b['reported_state']
    assert next(x for x in before if x['field']=='sacks')['reported_state']=='ZERO'
    assert any(x['dependency']=='measured_Lab_observations' and x['classification']=='NOT_VERIFIABLE' for x in r['unresolved_dependencies'])
    wrong=deepcopy(op);wrong['measurements']={'Main':{}}
    assert g.inspect_reference(path,wrong,'Main')['status']=='INPUT_REJECTED'
