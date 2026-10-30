"""C05 source-rendered structural controls, not App goldens or full-product UAT."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile

from docx import Document
import pytest

from c05_fluids import compare_document
from comparison_records import ROOT
from engineering_source import request


def operator(mode='Barite'):
    return {'job_type': 'CMT PLUG', 'fluids': {
        'Main': {'volume': 20, 'density': 118, 'pump_rate': '2-4'},
        'Spacer': {'volume': 20, 'density': 110, 'pump_rate': '2'}},
        'standalone_spacer': {'mode': mode, 'dead_vol': 5, 'tank_name': 'Mud Reserve Tanks',
            'limestone_available_lb': 1500, 'barite_available_lb': 10000}}


@pytest.fixture(scope='module')
def case():
    with tempfile.TemporaryDirectory(dir=ROOT/'tmp', prefix='c05-') as folder:
        path = Path(folder)/'synthetic.docx'
        op = operator()
        source = request('render_fluids', op, docx=str(path))
        yield path, op, source


def mutate(case, tmp_path, change):
    path, op, source = case
    # Inputs must stay under provisioned runtime, independent of pytest defaults.
    target = path.parent / (tmp_path.name + '.docx')
    document = Document(path)
    change(document)
    document.save(target)
    return compare_document(target, op)


def select(result, field, **identity):
    return [f for f in result['findings'] if f['field'] == field and all(f.get(k) == v for k, v in identity.items())]


def test_source_parity_order_and_independent_time_oracle(case):
    path, op, source = case
    ref = source['reference']
    # Independent, hand-specified schedule: two 20-bbl fluids at 2 bbl/min.
    assert ref['order'] == ['Spacer', 'Main']
    assert ref['total_pump_time_min'] == 20
    assert ref['fluids']['Main']['duration_str'] == '00:10'
    assert ref['fluids']['Main']['cumul_time_str'] == '00:20'
    assert 'project_state.py:refresh_fluids' in source['source_calls']
    assert 'standalone_spacer.py:calculate_spacer' in source['source_calls']
    result = compare_document(path, op)
    assert result['status'] == 'PARTIAL_FINDINGS_ONLY'
    assert result['findings'] and all(f['result'] == 'MATCH' for f in result['findings'])
    assert result['unverified_record_orders']  # unsupported prose/Lab never disappears


@pytest.mark.parametrize('mode,density', [('Limestone',90), ('Barite',110), ('Ferro Bar',140), ('Limestone + Barite',110)])
def test_all_configured_spacer_modes_and_display(case, mode, density):
    path, _, _ = case
    op = operator(mode); op['fluids']['Spacer']['density'] = density
    target = path.parent / (mode.replace(' ', '_')+'.docx')
    source = request('render_fluids', op, docx=str(target))
    findings = compare_document(target, op)['findings']
    assert all(f['result'] == 'MATCH' for f in findings)
    for f in findings:
        if f['field'] in ('spacer.base', 'spacer.gross'):
            assert f['locator'] and f['source_document_sha256'] and f['full_precision_expected'] is not None
    if mode == 'Limestone + Barite':
        assert source['reference']['spacer']['recipe']['inventory']['Limestone']['remaining_lb'] >= 0


def test_no_mandatory_inactive_fluids_and_range_time(case):
    path, _, _ = case
    op = {'job_type':'CSG 20"', 'fluids': {'Main':{'volume':20,'density':118,'pump_rate':'2-4'}}}
    target = path.parent/'only-main.docx'
    request('render_fluids', op, docx=str(target))
    result = compare_document(target, op)
    assert all(f['result'] == 'MATCH' for f in result['findings'])
    assert set(result['reference']['reference']['fluids']) == {'Main'}


@pytest.mark.parametrize('name', ['Pre Flush', 'Spacer Ahead', 'Spacer Behind', 'Displacement Fluid', 'Diesel'])
def test_optional_source_fluids(case, name):
    path, _, _ = case
    op = {'job_type':'CSG 20"','mud_type':'OBM', 'fluids': {name:{'volume':20,'density':52 if name=='Diesel' else 70,'pump_rate':2,
        'material_name':'Salt Saturated Water' if name=='Pre Flush' else 'Operator Fluid'}}}
    target = path.parent/(name.replace(' ','_')+'.docx')
    request('render_fluids', op, docx=str(target))
    assert all(f['result'] == 'MATCH' for f in compare_document(target,op)['findings'])


@pytest.mark.parametrize('mutation', ['volume', 'time', 'gross', 'tank', 'duplicate', 'wrong-header', 'split-paragraph', 'cross-fluid'])
def test_negative_mutations_fail_closed(case, tmp_path, mutation):
    def change(d):
        fluid = next(t for t in d.tables if t.cell(0,0).text == 'Fluids')
        spacer = next(t for t in d.tables if t.cell(0,0).text == 'Spacer Additives Data')
        if mutation == 'volume': fluid.cell(2,3).text = '999'
        if mutation == 'time': fluid.cell(2,5).text = '99:59'
        if mutation == 'gross': spacer.cell(3,5).text = '999999 lb'
        if mutation == 'tank':
            next(p for p in d.paragraphs if 'Mix above Additives' in p.text).text = 'NOTE 5: Mix above Additives in 20.0 bbl Fresh Water at Wrong Tank.'
        if mutation == 'duplicate': spacer._tbl.addnext(deepcopy(spacer._tbl))
        if mutation == 'wrong-header': spacer.cell(2,5).text = 'lbs or gal'
        if mutation == 'split-paragraph': fluid.cell(2,3).add_paragraph('20')
        if mutation == 'cross-fluid': fluid.cell(2,0).text = 'Main'
    result = mutate(case, tmp_path, change)
    assert any(f['result'] != 'MATCH' for f in result['findings'])
    if mutation in ('duplicate','wrong-header'):
        assert all(f['result'] == 'NOT_VERIFIABLE' for f in result['findings'] if f['field'].startswith('spacer.'))


@pytest.mark.parametrize('raw,state', [('', 'MISSING'), ('0','ZERO'), ('-','DASH'), ('N/A','NA_PRINTED'), ('bad','MALFORMED_OR_UNIT_UNPROVEN')])
def test_missing_zero_na_and_malformed_remain_distinct(case, tmp_path, raw, state):
    def change(d):
        next(t for t in d.tables if t.cell(0,0).text=='Fluids').cell(2,3).text=raw
    result = mutate(case,tmp_path,change)
    f = select(result,'fluid.volume',fluid='Spacer')[0]
    assert f['reported_state'] == state
    assert f['result'] != 'MATCH'


@pytest.mark.parametrize('bad', ['density','pump_rate','volume','tank','inventory','diesel','ct','inactive'])
def test_original_source_validation_rejects_invalid_reference(bad):
    op = operator()
    if bad in ('density','pump_rate','volume'): op['fluids']['Spacer'][bad] = 0
    if bad == 'tank': op['standalone_spacer']['tank_name'] = 'Unknown'
    if bad == 'inventory': op['standalone_spacer'].update(mode='Limestone + Barite',barite_available_lb=1)
    if bad == 'diesel': op['fluids']['Diesel']={'volume':1,'density':52,'pump_rate':1};op['mud_type']='WBM'
    if bad == 'ct': op['job_type']='CMT PLUG WITH CT';op['fluids']['Lead']=op['fluids'].pop('Main')
    if bad == 'inactive': op['fluids'].pop('Spacer')
    with pytest.raises(ValueError): request('fluids',op)


def test_cli_determinism_original_bytes_and_source_drift(case):
    path, op, _ = case
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    assert compare_document(path,op) == compare_document(path,op)
    inputs = path.parent/'operator.json';inputs.write_text(json.dumps(op))
    process = subprocess.run([sys.executable,'-B',str(ROOT/'c05_fluids.py'),str(inputs),str(path)],capture_output=True,text=True,check=True)
    assert json.loads(process.stdout)['status'] == 'PARTIAL_FINDINGS_ONLY'
    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
    # No source mutation: a copied inventory mismatch is covered by existing lock controls.
    from c02_engine import verify_source
    verify_source()


def test_symlink_escape_and_no_network_calculations(case):
    path, op, _ = case
    with pytest.raises(ValueError,match='without symlinks'):
        request('render_fluids',op,docx=str(path.parent/'../escape.docx'))
    from tools.ci import source_proof
    assert source_proof()['verified_files'] == 143


def test_source_child_network_and_write_guards():
    script = """import engineering_source as s, socket, subprocess
s.install_audit()
# Trusted imports may query host metadata in a read-only system subprocess.
# A blanket Popen prohibition broke real hosted imports; required guards remain.
subprocess.run(['/usr/bin/uname', '-p'], check=True, capture_output=True)
for action in (lambda: socket.socket().connect(('127.0.0.1',9)),
               lambda: open(s.SNAPSHOT/'forbidden-write', 'w'),
               lambda: open(s.ROOT.parent/'forbidden-write', 'w')):
    try: action()
    except RuntimeError: pass
    else: raise AssertionError('Boundary bypass')
"""
    subprocess.run([sys.executable,'-B','-c',script],cwd=ROOT,check=True,capture_output=True,text=True)
