"""Admission controls only: no authentic Golden or field measurements claimed."""
import hashlib
from pathlib import Path
import json
import sys
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

from comparison_records import ROOT
from golden_admission import JOBS, PIPELINE, SOURCE_SHA, TEMPLATE_SHA256, case_record, inspect_external, generate, public_metadata
from golden_source import validate_project


def synthetic_document(path):
    doc=Document();doc.add_heading('Executive Summary',1)
    doc.add_paragraph('Synthetic structural narrative. Not an App production report.')
    doc.add_heading('Cementing Procedure',1);doc.add_paragraph('1. Synthetic instruction.')
    doc.save(path);return path


def claimed(path):
    return {'job_type':'CSG 20"','source_sha':SOURCE_SHA,'template_sha256':TEMPLATE_SHA256,
        'document_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'independent_input_sha256':'1'*64}


@pytest.fixture
def case():
    with tempfile.TemporaryDirectory(dir=ROOT/'tmp',prefix='g03-admission-') as folder:
        path=synthetic_document(Path(folder)/'synthetic.docx');yield path


def test_fifteen_canonical_cases_match_locked_source_and_no_golden():
    sys.path.insert(0,str(ROOT/'snapshot'))
    import materials_db
    assert list(JOBS)==materials_db.JOB_TYPES and len(set(JOBS))==15
    rows=[case_record(job) for job in JOBS]
    assert [r['case_id'] for r in rows]==[f'J{i:02}' for i in range(1,16)]
    assert all(r['golden_pass_allowed'] is False and r['acceptance_A_H']['F']=='NOT_RUN' and r['acceptance_A_H']['G']=='NOT_RUN' for r in rows)
    manifest=json.loads((ROOT/'docs/uat/job-manifest.json').read_text())
    assert [r['job_type'] for r in manifest['jobs']]==list(JOBS)
    assert all(r['status']=='UNAVAILABLE' and r['generation_status']=='NOT_RUN' for r in manifest['jobs'])


def test_synthetic_structure_and_correct_hash_claims_do_not_authenticate(case):
    before=case.read_bytes();r=inspect_external(case,claimed(case))
    assert r['status']=='GENERATION_BLOCKED' and r['blocker']=='FRESH_LEGITIMATE_PRODUCTION_EXECUTION_REQUIRED'
    assert r['golden_pass_allowed'] is False and case.read_bytes()==before


@pytest.mark.parametrize('key',['source_sha','template_sha256','document_sha256','independent_input_sha256'])
def test_wrong_pins_or_missing_independent_provenance_fail_closed(case,key):
    metadata=claimed(case);metadata[key]='bad'
    assert inspect_external(case,metadata)['status']=='INPUT_INCOMPLETE'


def test_self_asserted_execution_approvals_or_golden_pass_are_not_receipts(case):
    for extra in ({'pipeline_calls':list(PIPELINE)},{'authentic':True},{'reviewed':True},{'status':'GOLDEN_PASS'}):
        assert inspect_external(case,{**claimed(case),**extra})['status']=='INPUT_INCOMPLETE'


@pytest.mark.parametrize('mutation',['duplicate-summary','duplicate-procedure','numPr','field','hidden','damaged'])
def test_structural_trust_negative_controls(case,mutation):
    d=Document(case)
    if mutation.startswith('duplicate'):
        d.add_heading('Executive Summary' if mutation.endswith('summary') else 'Cementing Procedure',1)
    elif mutation=='numPr':d.paragraphs[1]._p.get_or_add_pPr().append(OxmlElement('w:numPr'))
    elif mutation=='field':d.paragraphs[3]._p.append(OxmlElement('w:fldSimple'))
    elif mutation=='hidden':d.paragraphs[1].runs[0]._r.get_or_add_rPr().append(OxmlElement('w:vanish'))
    if mutation=='damaged':case.write_bytes(b'not a Word package')
    else:d.save(case)
    assert inspect_external(case,claimed(case))['status']=='INPUT_INCOMPLETE'


@pytest.mark.parametrize('forged',['reviewed','review_signature','lab_source_signatures','calculations_prepared','_compiled_doc_bytes'])
def test_input_cannot_supply_cached_or_fabricated_review_state(forged):
    p={'job_type':'CSG 20"','lab_qc_params':{'Main':{forged:True}}}
    with pytest.raises(ValueError):validate_project(p)


def test_public_projection_never_exports_private_data_or_promotes():
    record=case_record('CSG 20"');record.update(status='GENERATION_BLOCKED',private_gate_details='private values',
        private_document_path='/private/client.docx',operator={'well_name':'private'},source_document_sha256='a'*64)
    public=public_metadata(record)
    assert 'private' not in json.dumps(public) and 'operator' not in public
    for status in ('GOLDEN_PASS','PARTIAL_VERIFIED'):
        with pytest.raises(ValueError):public_metadata({**record,'status':status})


def test_private_path_escape_and_unavailable_are_not_admitted(case):
    alias=case.parent/'escape.docx';alias.symlink_to(case)
    assert inspect_external(alias,claimed(case))['status']=='INPUT_INCOMPLETE'
    assert inspect_external(case.parent/'absent.docx',claimed(case))['status']=='UNAVAILABLE'


def test_real_default_export_gate_blocks_without_lab_data_and_review():
    reference={'kind':'SYNTHETIC_TEST_INPUT','project':{'job_type':'CSG 20"',
        'fluids_config':{'active':['Main'],'params':{'Main':{'volume':50,'density':104,'pump_rate':2}}},
        'cement_params':{'Main':{'auto_calc':True,'base_cement':'Cement G Delijan','cmt_sg':3.2,'cmt_sg_source':'manual','dead_vol':0,'tank_name':'Mud Reserve Tanks'}},
        'cement_additives_dfs':{'Main':[]},'well_geometry':{'version':1,'type':'Vertical','td_m':1000},
        'geothermal_config':{'version':1,'source':'BHST','surface_temp':80,'value':180},'bhct':150,'mud_density':75}}
    r=generate(reference)
    assert r['status']=='GENERATION_BLOCKED' and r['source_document_sha256'] is None
    assert r['last_stage']=='build_master_context' and 'phase_10_procedure.py:build_master_context' in r['pipeline_calls']
    assert 'Phase VII' in r['private_gate_details'] and r['source_files_verified']==143
    assert 'well_report.py:postprocess_well_report' not in r['pipeline_calls']
    assert public_metadata(r)['hosted_golden_uat']=='NOT_RUN_HOSTED'
