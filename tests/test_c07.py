"""Independent synthetic observations and genuine locked-source Lab projections."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile

from docx import Document
from docx.table import Table
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from decimal import Decimal
import pytest

from checker_parser import parse_docx
from comparison_records import ROOT, finish
from engineering_source import request
from c07_lab import compare_document
from checker_parser import norm


def operator():
    dials=[[168,116,64,41,24,8,6],[175,124,68,45,29,12,9],[150,108,70,52,38,18,14]]
    rheology={key:{'selected':True,'readings':dict(zip(map(str,(300,200,100,60,30,6,3)),values)),
        'gel_10_sec':8+i,'gel_10_min':12+i} for i,(key,values) in enumerate(zip(('surface_down','bhct_up','bhct_down'),dials))}
    qc={'rheology':rheology,'free_water':0.5,'free_water_45':0.25,'surface_hardened_hours':8,'api_fl':6.2,
        'thickening_test_start':'08:15','thickening_cell':2,'thickening_30_bc':'01:00','thickening_50_bc':'02:00','thickening_time':'03:00',
        'compressive':{'uca':{'selected':True,'cs_8':1000,'cs_12':2000,'cs_24':3000},
            'crush':{'selected':True,**{'force_'+str(n):1000 for n in range(1,5)}}}}
    cement={'job_type':'CSG 20"','active_fluids':['Main'],'slurries':{'Main':{
        'client_volume_bbl':50,'density_pcf':104,'base_cement':'Cement G Delijan','cement_sg':3.2,'additives':[]}}}
    return {'job_type':cement['job_type'],'active_fluids':['Main'],'conditions':{'bhct':150,'bhst':180,'bhsp':1000.5},
        'confirmed':True,'measurements':{'Main':qc},'cement_operator':cement}


@pytest.fixture(scope='module')
def case():
    with tempfile.TemporaryDirectory(dir=ROOT/'tmp',prefix='c07-') as folder:
        p=Path(folder)/'synthetic.docx';op=operator()
        source=request('render_lab',op,docx=str(p))
        yield p,op,source


def test_original_engine_oracles_and_supported_word_claims(case):
    p,op,source=case
    ref=source['reference']['fluids']['Main']
    # Independent corpus A expected display (locked upstream test_uat_round5_rheology).
    assert ref['display']['rheology']['surface_down']['pv_cp']=='159.218'
    assert ref['display']['rheology']['surface_down']['ty_lbf_100ft2']=='9.47'
    assert ref['lab']['compressive']['crush']['result']==250
    assert ref['lab']['compressive']['uca']['result']==3000
    assert ref['lab']['api_fl_collected']==6.2 and ref['lab']['fluid_loss']==12.4
    assert ref['payload_admitted'] is True
    assert {'rheology.py:calculate_cemcade_bingham','project_state.py:refresh_lab_payloads',
        'engineering_tools.py:validate_compressive_test','engineering_tools.py:validate_thickening_test'}<=set(source['source_calls'])
    result=compare_document(p,op)
    assert result['status']=='PARTIAL_FINDINGS_ONLY' and result['findings']
    assert all(f['result']=='MATCH' for f in result['findings'])
    assert field(result,'lab.rheology.surface_down.pv_cp')['reference_kind']=='SOURCE_DERIVED_FROM_INDEPENDENT_MEASUREMENTS'
    assert field(result,'lab.free_water')['reference_kind']=='INDEPENDENT_MEASUREMENT'
    assert field(result,'lab.recipe_mass',material='Cement G Delijan')['reference_kind']=='AUTOMATIC_SOURCE_FORMULATION'
    assert result['unverified_record_orders']
    assert all('iod' not in f['field'] for f in result['findings'])


def test_bound_unverified_records_remain_explicit(case):
    p,_,_=case;parsed=parse_docx(p)
    record=next(r for r in parsed['records'] if r['raw_text']=='12.4')
    result=finish(parsed,{},[{'locator':record['locator'],'result':'NOT_VERIFIABLE'}],[])
    assert record['source_order'] in result['unverified_record_orders']


def field(result,name,fluid='Main',material=None):
    return next(f for f in result['findings'] if f['field']==name and f['slurry']==fluid and (material is None or f.get('material','').casefold()==material.casefold()))


def cell_for(document,finding):
    loc=finding['locator'];assert loc and not loc.get('table_path')
    return Table(document.element.body[loc['body_block']],document._body).cell(loc['row'],loc['column'])


@pytest.mark.parametrize('mutation',['pv','cross-dataset','unit','span','duplicate','split','cross-slurry','filtrate','crush','condition'])
def test_word_mutations_never_falsely_match(case,monkeypatch,tmp_path,mutation):
    p,op,source=case;monkeypatch.setattr('c07_lab.request',lambda *a,**k:source)
    baseline=compare_document(p,op);d=Document(p)
    pv=field(baseline,'lab.rheology.surface_down.pv_cp');cell=cell_for(d,pv)
    table=next(t for t in d.tables if t.cell(0,0).text.strip()=='Rheology Test')
    if mutation=='pv':cell.text='999.999'
    if mutation=='cross-dataset':
        other=cell_for(d,field(baseline,'lab.rheology.bhct_down.pv_cp'));a,b=cell.text,other.text;cell.text=b;other.text=a
    if mutation=='unit':next(r for r in table.rows if norm(r.cells[0].text)=='pv (cp)').cells[0].text='Pv (Pa s)'
    if mutation=='span':
        span=OxmlElement('w:gridSpan');span.set(qn('w:val'),'2');cell._tc.get_or_add_tcPr().append(span)
    if mutation=='duplicate':table._tbl.addnext(deepcopy(table._tbl))
    if mutation=='split':cell.add_paragraph('159.218')
    if mutation=='cross-slurry':
        heading=d.add_paragraph('Tail Cement Slurry Results');table._tbl.addprevious(heading._p)
    if mutation=='filtrate':cell_for(d,field(baseline,'lab.filtrate_30min')).text='99.0 ml'
    if mutation=='crush':cell_for(d,field(baseline,'lab.crush.force_1')).text='9999'
    if mutation=='condition':cell_for(d,field(baseline,'lab.condition.bhct')).text='999 degF'
    target=p.parent/(tmp_path.name+'.docx');d.save(target);result=compare_document(target,op)
    if mutation in ('pv','cross-dataset'):assert field(result,'lab.rheology.surface_down.pv_cp')['result']=='DISCREPANCY'
    elif mutation in ('unit','span','duplicate','split','cross-slurry'):
        assert all(f['result']=='NOT_VERIFIABLE' for f in result['findings'] if f['field'].startswith('lab.rheology.'))
    elif mutation=='filtrate':assert field(result,'lab.filtrate_30min')['result']=='DISCREPANCY'
    elif mutation=='crush':assert field(result,'lab.crush.force_1')['result']=='DISCREPANCY'
    else:
        assert field(result,'lab.thickening_time')['result']=='NOT_VERIFIABLE'
        assert field(result,'lab.rheology.bhct_down.pv_cp')['result']=='NOT_VERIFIABLE'
    assert result['status']=='PARTIAL_FINDINGS_ONLY'


@pytest.mark.parametrize('raw,state',[('', 'MISSING'),('0','ZERO'),('-','DASH'),('N/A','NA_PRINTED'),('bad','MALFORMED_OR_UNIT_UNPROVEN')])
def test_observed_reading_states_remain_distinct(case,monkeypatch,tmp_path,raw,state):
    p,op,source=case;monkeypatch.setattr('c07_lab.request',lambda *a,**k:source)
    baseline=compare_document(p,op);d=Document(p)
    cell_for(d,field(baseline,'lab.rheology.surface_down.300')).text=raw
    target=p.parent/(tmp_path.name+'.docx');d.save(target)
    f=field(compare_document(target,op),'lab.rheology.surface_down.300')
    assert f['reported_state']==state and f['result']!='MATCH'


@pytest.mark.parametrize('bad',['rpm','too-few-readings','free-water','filtrate','missing-hours','zero-time','over-time','compressive','cached-fit','inactive','job','sg','dv','confirmation'])
def test_source_validation_and_independent_input_fail_closed(bad):
    op=operator();qc=op['measurements']['Main']
    if bad=='rpm':qc['rheology']['surface_down']['readings']['600']=200
    if bad=='too-few-readings':qc['rheology']['surface_down']['readings']={'300':168}
    if bad=='free-water':qc['free_water']=-1
    if bad=='filtrate':qc['api_fl']=-1
    if bad=='missing-hours':qc.pop('surface_hardened_hours')
    if bad=='zero-time':qc['thickening_time']='00:00'
    if bad=='over-time':qc['thickening_time']='24:01'
    if bad=='compressive':qc['compressive']['uca'].pop('cs_12')
    if bad=='cached-fit':qc['pv_cp']=159.218
    if bad=='inactive':op['measurements']['Tail']=deepcopy(qc)
    if bad=='job':op['cement_operator']['job_type']='CMT PLUG'
    if bad=='sg':op['cement_operator']['slurries']['Main'].pop('cement_sg')
    if bad=='dv':op['corrected_dead_volume_bbl']=5
    if bad=='confirmation':op['confirmed']='yes'
    with pytest.raises(ValueError):request('lab',op)


@pytest.mark.parametrize('bad',[float('nan'),float('inf'),True,None])
def test_non_measurements_cannot_become_zero(bad):
    op=operator();op['measurements']['Main']['api_fl']=bad
    with pytest.raises(ValueError):request('lab',op)


def test_unconfirmed_and_missing_filtrate_have_no_derived_prediction(case,monkeypatch):
    p,op,_=case
    for omit in (False,True):
        independent=deepcopy(op);independent['confirmed']=False
        if omit:independent['measurements']['Main'].pop('api_fl')
        source=request('lab',independent);assert source['reference']['fluids']['Main']['payload_admitted'] is False
        monkeypatch.setattr('c07_lab.request',lambda *a,**k:source)
        result=compare_document(p,independent);f=field(result,'lab.api_fluid_loss')
        assert f['result']=='NOT_VERIFIABLE' and f['full_precision_expected'] is None
        assert f['raw_reported_text']=='12.4'
        loc=f['locator'];assert next(r['source_order'] for r in result['records'] if r['locator']==loc) in result['unverified_record_orders']
        assert field(result,'lab.filtrate_30min')['result']==('NOT_VERIFIABLE' if omit else 'MATCH')


def test_surface_only_and_partial_rpm_do_not_invent_other_observations(case,monkeypatch):
    p,op,_=case
    surface=deepcopy(op['measurements']['Main']['rheology']['surface_down'])
    surface['readings'].pop('300')  # original source permits an omitted point, not a zero observation
    independent={'job_type':op['job_type'],'active_fluids':['Main'],'measurements':{'Main':{'rheology':{'surface_down':surface}}}}
    source=request('lab',independent);monkeypatch.setattr('c07_lab.request',lambda *a,**k:source)
    result=compare_document(p,independent)
    assert field(result,'lab.rheology.surface_down.300')['result']=='NOT_VERIFIABLE'
    assert field(result,'lab.rheology.surface_down.300')['full_precision_expected'] is None
    assert not any('bhct_down' in f['field'] for f in result['findings'])
    assert field(result,'lab.api_fluid_loss')['result']=='NOT_VERIFIABLE'


@pytest.mark.parametrize('fluid,job,mud',[('Scavenger','CMT PLUG','WBM'),('Diesel','CMT PLUG','OBM'),('Pre Flush','CMT PLUG','WBM'),('Main','CMT PLUG WITH CT','WBM'),('Spacer Ahead','CMT PLUG','WBM')])
def test_unsupported_word_roles_do_not_borrow_main_lab(case,fluid,job,mud):
    p,base,_=case
    qc={'rheology':{'bhct_down':deepcopy(base['measurements']['Main']['rheology']['bhct_down'])}}
    op={'job_type':job,'mud_type':mud,'active_fluids':[fluid],'measurements':{fluid:qc},'conditions':base['conditions']}
    result=compare_document(p,op)
    assert result['status']=='PARTIAL_FINDINGS_ONLY'
    assert result['findings'] and all(f['result']=='NOT_VERIFIABLE' for f in result['findings'])


def test_two_slurries_do_not_cross_readings_or_recipe_masses(case):
    p,base,_=case;op=deepcopy(base)
    op['active_fluids']=['Lead','Tail'];op['measurements']={'Lead':op['measurements']['Main'],'Tail':deepcopy(op['measurements']['Main'])}
    op['measurements']['Tail']['rheology']['surface_down']=deepcopy(op['measurements']['Lead']['rheology']['bhct_up'])
    op['measurements']['Tail']['api_fl']=8
    cop=op['cement_operator'];cop['active_fluids']=op['active_fluids']
    original=cop['slurries'].pop('Main');cop['slurries']={'Lead':original,'Tail':dict(original,density_pcf=118,client_volume_bbl=20)}
    target=p.parent/'two-slurries.docx';request('render_lab',op,docx=str(target))
    result=compare_document(target,op)
    assert all(f['result']=='MATCH' for f in result['findings'])
    assert field(result,'lab.rheology.surface_down.pv_cp','Lead')['source_rounded_expected']=='159.218'
    assert field(result,'lab.rheology.surface_down.pv_cp','Tail')['source_rounded_expected']=='163.158'
    assert field(result,'lab.api_fluid_loss','Tail')['full_precision_expected']==16


@pytest.mark.parametrize('dose,state,tolerance',[(0.009,'Powder','0'),(0.01,'Powder','.3'),(1,'Powder','.3'),(1.001,'Powder','1'),(0.2,'Liquid','.3')])
def test_recipe_mass_owner_threshold_edges_and_outside(case,monkeypatch,dose,state,tolerance):
    p,op,_=case;op=deepcopy(op)
    additive={'Material Type':'Custom '+state,'Name':'C07 Additive','Physical State':state,'Mix Method':'In Mix Water','User Input':dose,'Density':1.4 if state=='Powder' else 9.0}
    op['cement_operator']['slurries']['Main']['additives']=[additive]
    target=p.parent/('recipe-'+str(dose)+'-'+state+'.docx');source=request('render_lab',op,docx=str(target))
    monkeypatch.setattr('c07_lab.request',lambda *a,**k:source)
    baseline=compare_document(target,op);f=field(baseline,'lab.recipe_mass',material='C07 Additive')
    assert f['result']=='MATCH' and Decimal(f['tolerance'])==Decimal(tolerance)
    for outside in (False,True):
        d=Document(target);delta=Decimal(tolerance)+(Decimal('.01') if outside else 0)
        cell_for(d,f).text=str(Decimal(f['source_rounded_expected'])+delta)
        changed=p.parent/('edge-'+str(dose)+'-'+state+'-'+str(outside)+'.docx');d.save(changed)
        got=field(compare_document(changed,op),'lab.recipe_mass',material='C07 Additive')
        assert got['result']==('DISCREPANCY' if outside else 'MATCH')


@pytest.mark.parametrize('mode',['Barite','Limestone + Barite'])
def test_configured_spacer_recipe_and_measured_rheology(case,mode):
    p,base,_=case;op=deepcopy(base)
    op['active_fluids'].append('Spacer');op['cement_operator']['active_fluids'].append('Spacer')
    op['measurements']['Spacer']={'rheology':deepcopy(op['measurements']['Main']['rheology'])}
    op['spacer_operator']={'job_type':op['job_type'],'fluids':{'Spacer':{'volume':20,'density':110,'pump_rate':2}},
        'standalone_spacer':{'mode':mode,'dead_vol':5,'tank_name':'Mud Reserve Tanks','limestone_available_lb':1500,'barite_available_lb':10000}}
    target=p.parent/('spacer-'+mode.replace(' ','_')+'.docx');request('render_lab',op,docx=str(target))
    result=compare_document(target,op);spacer=[f for f in result['findings'] if f['slurry']=='Spacer']
    assert spacer and all(f['result']=='MATCH' for f in spacer)
    assert field(result,'lab.recipe_mass','Spacer','Fresh Water')['unit']=='g/600 mL'


def test_client_volume_changes_neither_lab_recipe_nor_measured_results(case):
    _,op,source=case;op=deepcopy(op);op['cement_operator']['slurries']['Main']['client_volume_bbl']=100
    changed=request('lab',op)['reference'];before=source['reference']
    assert changed['cement_reference']['slurries']['Main']['raw_lab']==before['cement_reference']['slurries']['Main']['raw_lab']
    assert changed['fluids']['Main']['rheology']==before['fluids']['Main']['rheology']
    assert changed['fluids']['Main']['lab']['fluid_loss']==12.4


def test_duplicate_materials_remain_ambiguous(case,monkeypatch):
    p,op,_=case;op=deepcopy(op)
    row={'Material Type':'Custom Powder','Name':'C07 Duplicate','Physical State':'Powder','Mix Method':'In Mix Water','User Input':1,'Density':1.4}
    op['cement_operator']['slurries']['Main']['additives']=[row,dict(row,**{'User Input':2})]
    target=p.parent/'duplicates.docx';source=request('render_lab',op,docx=str(target));monkeypatch.setattr('c07_lab.request',lambda *a,**k:source)
    findings=[f for f in compare_document(target,op)['findings'] if f.get('material')=='C07 Duplicate']
    assert findings and all(f['result']=='NOT_VERIFIABLE' for f in findings)


def test_source_conditions_thickening_endpoint_and_no_fabrication(case):
    _,op,_=case;op=deepcopy(op);op['measurements']['Main']['thickening_time']='24:00'
    assert request('lab',op)['reference']['fluids']['Main']['lab']['thickening_time']=='24:00'
    op['conditions']['bhct']=op['conditions']['bhst']+1
    with pytest.raises(ValueError,match='BHCT/BHST'):request('lab',op)


def test_cli_determinism_original_bytes_and_identity_guards(case):
    p,op,_=case;digest=hashlib.sha256(p.read_bytes()).hexdigest();inputs=p.parent/'operator.json';inputs.write_text(json.dumps(op))
    before=inputs.read_bytes()
    values=[subprocess.check_output([sys.executable,'-B',str(ROOT/'c07_lab.py'),str(inputs),str(p)],text=True) for _ in range(2)]
    assert values[0]==values[1] and json.loads(values[0])['status']=='PARTIAL_FINDINGS_ONLY'
    assert inputs.read_bytes()==before and hashlib.sha256(p.read_bytes()).hexdigest()==digest
    assert json.loads(values[0])['reference']['operator_input_sha256']==hashlib.sha256(json.dumps(op,sort_keys=True,allow_nan=False).encode()).hexdigest()
    from tools.provision import verify_candidate
    from tools.ci import source_proof
    verify_candidate(ROOT);assert source_proof()['verified_files']==143
