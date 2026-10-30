"""Exact source narratives and independent text/order mutations; synthetic, not goldens."""
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import subprocess
import sys
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

from checker_parser import parse_docx
from comparison_records import ROOT
from narrative_source import request
from c08_summary import compare_document, equivalent


def operator(job='CMT PLUG'):
    return {'job_type':job,'geo_md':1000,
        'hardware':[{'Description':'Casing','MD (m)':'1000','Size (in)':'7','ID (in)':6},
            {'Description':'Previous Casing','MD (m)':'500','Size (in)':'9.625','ID (in)':8.535},
            {'Description':'Open Hole Size','MD (m)':'500-1000','Size (in)':'8.5','ID (in)':8.5}],
        'placement_config':{'job_type':job,'target_row':'__manual__','manual_depth_m':1000,'chart_excess_mode':'ignore'},
        'fluids':{'Main':{'volume':50,'density':104,'pump_rate':2}},
        'cement_operator':{'job_type':job,'active_fluids':['Main'],'slurries':{'Main':{
            'client_volume_bbl':50,'density_pcf':104,'base_cement':'Cement G Delijan','cement_sg':3.2,'additives':[]}}},
        'tops':{'Main':{'top_job_type':job,'top_mode':'Depth (m MD)','top_depth':900}},
        'executive_summary_config':{'version':1,'job_type':job,'placement_context_rows':[],
            'per_slurry':{'Main':{'volume_basis':'client_request'}}}}


@pytest.fixture(scope='module')
def case():
    with tempfile.TemporaryDirectory(dir=ROOT/'tmp',prefix='c08-') as folder:
        p=Path(folder)/'synthetic.docx';op=operator()
        source=request(op,'render_summary',str(p))
        yield p,op,source


def summary_paragraph(document):
    return next(p for p in document.paragraphs if p.text.startswith('Enclosed are our recommendations'))


def test_original_generator_real_session_and_hand_checked_oracle(case):
    p,op,source=case;ref=source['reference'];text=ref['text']
    assert text.split('\n\n')[2]==('Cement plug will be set at 1000.0 m MD. Neat cement design will be used to cement inside casing / open hole; '
        '104 pcf main cement slurry will be cemented from 1000.0 m to 900.0 m. '
        'The calculated cement volume is based on client request, giving 50.0 bbl for main slurry.')
    assert text.startswith('Enclosed are our recommendations for NIDC Cement Engineering and Planning Department intervention on the referenced well.')
    assert 'A pre-job safety meeting will be held with client representatives' in text
    assert ref['session_state_implementation']=='streamlit.runtime.state.session_state_proxy'
    assert ref['approval_flags_supplied'] is False
    assert 'phase_10_procedure.py:generate_executive_summary' in ref['source_calls']
    assert source['source_module_sha256']['phase_10_procedure.py']
    result=compare_document(p,op)
    assert result['status']=='PARTIAL_FINDINGS_ONLY' and result['findings'][0]['result']=='MATCH'
    assert result['findings'][0]['raw_reported_text']==text
    assert result['findings'][0]['source_document_sha256']==sha256(p.read_bytes()).hexdigest()
    assert result['unverified_record_orders'] and result['records']==parse_docx(p)['records']


def test_ordinary_casing_and_client_volume_use_original_placement():
    op=operator('CSG 7"');source=request(op)
    assert source['reference']['text'].split('\n\n')[2].startswith('7" Casing will be set and cemented at 1000.0 m MD.')
    assert 'giving 50.0 bbl for main slurry.' in source['reference']['text']
    assert 'cement_chart_model.py:resolve_cement_placement' in source['reference']['source_calls']


@pytest.mark.parametrize('mutation',['safety','depth','volume','density','fluid','omit','reorder','punctuation','number','case'])
def test_meaning_mutations_are_discrepancies(case,monkeypatch,mutation):
    p,op,source=case;monkeypatch.setattr('c08_summary.request',lambda *a:source)
    d=Document(p);para=summary_paragraph(d);text=para.text
    replacements={'safety':('must adhere','may ignore'),'depth':('1000.0','999.0'),
        'volume':('50.0 bbl','51.0 bbl'),'density':('104 pcf','118 pcf'),
        'fluid':('main cement','tail cement'),'punctuation':('safe operation.','safe operation?'),
        'number':('1000.0 m MD','10000.0 m MD'),'case':('NIDC','nidc')}
    if mutation=='omit':text=text.replace(text.split('\n\n')[1]+'\n\n','')
    elif mutation=='reorder':parts=text.split('\n\n');text='\n\n'.join([parts[1],parts[0],parts[2]])
    else:a,b=replacements[mutation];assert a in text;text=text.replace(a,b)
    para.text=text;target=p.parent/(mutation+'.docx');d.save(target)
    result=compare_document(target,op)
    assert result['findings'][0]['result']=='DISCREPANCY'
    assert result['records']==parse_docx(target)['records']


@pytest.mark.parametrize('mutation',['repeat','split-physical','duplicate-heading','unknown-heading','table','wrong-section'])
def test_structure_ambiguity_never_matches(case,monkeypatch,mutation):
    p,op,source=case;monkeypatch.setattr('c08_summary.request',lambda *a:source)
    d=Document(p);para=summary_paragraph(d)
    if mutation=='repeat':para._p.addnext(deepcopy(para._p))
    elif mutation=='split-physical':
        text=para.text.split('\n\n');para.text=text[0];new=deepcopy(para._p);para._p.addnext(new)
        from docx.text.paragraph import Paragraph
        Paragraph(new,para._parent).text='\n\n'.join(text[1:])
    elif mutation in ('duplicate-heading','unknown-heading'):
        heading=next(p for p in d.paragraphs if 'Executive Summary' in p.text)
        new=deepcopy(heading._p);para._p.addnext(new)
        from docx.text.paragraph import Paragraph
        if mutation=='unknown-heading':Paragraph(new,para._parent).text='Unknown narrative subsection'
    elif mutation=='table':
        table=d.add_table(rows=1,cols=1);table.cell(0,0).text='Additional unproved narrative';para._p.addnext(table._tbl)
    else:next(p for p in d.paragraphs if 'Executive Summary' in p.text).text='Well Data Summary'
    target=p.parent/(mutation+'.docx');d.save(target)
    assert compare_document(target,op)['findings'][0]['result']=='NOT_VERIFIABLE'


def test_split_runs_keep_full_text_and_original_bytes(case,monkeypatch):
    p,op,source=case;monkeypatch.setattr('c08_summary.request',lambda *a:source)
    d=Document(p);para=summary_paragraph(d);text=para.text;para.clear()
    for piece in (text[:7],text[7:139],text[139:]):para.add_run(piece)
    target=p.parent/'runs.docx';d.save(target);before=target.read_bytes()
    assert compare_document(target,op)['findings'][0]['result']=='MATCH'
    assert target.read_bytes()==before
    assert equivalent('a\n\nb')!=equivalent('a b') and equivalent('A!')!=equivalent('a!')


@pytest.mark.parametrize('missing',['geo_md','hardware','tops','executive_summary_config','cement_operator'])
def test_incomplete_independent_inputs_fail_closed(case,missing):
    p,op,_=case;op=deepcopy(op);del op[missing]
    result=compare_document(p,op)
    assert result['findings'][0]['result']=='NOT_VERIFIABLE'
    assert len(result['unverified_record_orders'])==len(result['records'])


@pytest.mark.parametrize('bad',['ct','approval','crossed','metadata','top','context'])
def test_invalid_or_unsupported_source_context(case,bad):
    _,op,_=case;op=deepcopy(op)
    if bad=='ct':op['job_type']='CMT PLUG WITH CT'
    elif bad=='approval':op['calculations_prepared']=True
    elif bad=='crossed':op['fluids']['Main']['volume']=51
    elif bad=='metadata':op['executive_summary_config']['job_type']='CMT SQUEEZE'
    elif bad=='top':op['tops']['Main']['top_depth']=1001
    else:op['executive_summary_config']['placement_context_rows']=['unknown']
    with pytest.raises(ValueError):request(op)


def test_independent_reviewed_manual_text_and_hash(case):
    p,_,source=case;text=source['reference']['text']
    op={'job_type':'CMT PLUG','manual_summary':{'text':text,'approval_reference':'SYNTHETIC independent review record example',
        'text_sha256':sha256(text.encode()).hexdigest()}}
    result=compare_document(p,op)
    assert result['findings'][0]['result']=='MATCH'
    assert result['findings'][0]['reference_kind']=='INDEPENDENT_APPROVED_MANUAL_TEXT'
    op['manual_summary']['approval_reference']=''
    assert compare_document(p,op)['findings'][0]['result']=='NOT_VERIFIABLE'
    op['manual_summary']['approval_reference']='Independent record';op['manual_summary']['text_sha256']='0'*64
    assert compare_document(p,op)['findings'][0]['result']=='NOT_VERIFIABLE'


def test_cli_json_determinism_and_docx_input_unchanged(case):
    p,op,_=case;operator_path=p.parent/'operator.json';operator_path.write_text(json.dumps(op))
    before=(p.read_bytes(),operator_path.read_bytes())
    command=[sys.executable,'-B',str(ROOT/'c08_summary.py'),str(p),str(operator_path)]
    a=subprocess.run(command,text=True,capture_output=True,check=True);b=subprocess.run(command,text=True,capture_output=True,check=True)
    assert a.stdout==b.stdout and json.loads(a.stdout)['findings'][0]['result']=='MATCH'
    assert before==(p.read_bytes(),operator_path.read_bytes())
