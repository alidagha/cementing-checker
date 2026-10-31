"""Bound selection and physical claims; private synthetic C04/source oracles."""
from copy import deepcopy
from functools import lru_cache
from hashlib import sha256
import json
from pathlib import Path
import subprocess
import sys
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
import pytest
import g06b_dead_volume as g


@pytest.fixture(scope='session')
def controls():
    with tempfile.TemporaryDirectory(prefix='g06b-', dir=g.runtime.HERE/'tmp') as folder:
        folder = Path(folder)
        op = json.loads((g.runtime.HERE/'fixtures/operator-main.json').read_text())
        main = folder/'main.docx'
        main.write_bytes((g.runtime.HERE/'fixtures/synthetic-main.docx').read_bytes())
        cases = {'main': {'path': main, 'op': op}}
        for name in ('dv0', 'dv10', 'zero-dose', 'two'):
            operator = deepcopy(op)
            if name == 'zero-dose':
                for row in operator['slurries']['Main']['additives']: row['User Input'] = 0
            if name == 'two':
                recipe = operator['slurries'].pop('Main')
                operator.update(job_type='CSG 13 3/8"', active_fluids=['Lead', 'Tail'])
                operator['slurries'] = {'Lead': deepcopy(recipe), 'Tail': dict(recipe, density_pcf=118, client_volume_bbl=30)}
            path = folder/(name+'.docx')
            expected = g.runtime.call({'action': 'fixture', 'operator': operator, 'docx': str(path),
                'fixture': {f: {'dead': 0 if name == 'dv0' else 10 if name == 'dv10' else 2 if f == 'Tail' else 5,
                    'tank': 'Batch Mixer Tanks' if f == 'Tail' else 'Mud Reserve Tanks'} for f in operator['slurries']}})
            cases[name] = {'path': path, 'op': operator, 'expected': expected}
        yield cases


@lru_cache(maxsize=16)
def inspected(path, operator, slurry):
    result = g.inspect_selection(path, json.loads(operator), slurry)
    assert result['status'] == 'PARTIAL_FINDINGS_ONLY', result
    return result


def request(case, dead=10, slurry='Main'):
    before = inspected(str(case['path']), json.dumps(case['op'], sort_keys=True), slurry)
    return {'schema': g.SCHEMA, 'bindings': deepcopy(before['bindings']), 'slurry': slurry,
        'corrected_dead_volume_bbl': dead}


def impact(case, dead=10, slurry='Main'):
    result = g.dead_volume_impact(case['path'], case['op'], request(case, dead, slurry))
    assert result['status'] == 'PARTIAL_FINDINGS_ONLY', result
    return result


def test_bound_impact_matches_independent_C04_and_source_rendered_oracle(controls):
    case = controls['main']; original = case['path'].read_bytes(); op = deepcopy(case['op'])
    r = impact(case); ledger = r['impact_ledger']
    oracle = g.preview_source.preview(case['path'], op, {'slurry': 'Main', 'corrected_dead_volume_bbl': 10})
    assert r['c04'] == oracle and r['original_c03'] == oracle['c03']
    assert ledger['source_after_tables'] == {k: controls['dv10']['expected']['fixture_expected']['Main'][k] for k in ('blend', 'additives')}
    assert r['bindings']['selected_claims_sha256'] == g.digest(r['selected_claims'])
    assert r['bindings']['original_document_sha256'] == sha256(original).hexdigest()
    assert r['bindings']['operator_sha256'] == g.digest(op)
    assert ledger['source_after_NOTE_display_bbl'] == '44.2'
    assert ledger['source_after_total_preparation_water_full_precision_bbl'] == oracle['selected_dead_volume_preview']['new_preparation_water_full_precision_bbl']
    assert len(ledger['gross_additives']) == 3
    assert all(x['original_Word_records'][0]['locator'] and x['original_Word_records'][0]['raw_text'] for x in ledger['gross_additives'])
    assert ledger['historical_exact_delta'] is None and ledger['historical_change'] == 'NOT_VERIFIABLE'
    assert all(x['source_after_full_precision'] is None and x['historical_exact_delta'] is None for x in ledger['gross_additives'])
    assert r['docx_written'] is r['word_modified'] is False and case['op'] == op and case['path'].read_bytes() == original
    assert {x['classification'] for x in r['unresolved_dependencies']} == {'MANUAL_REVIEW_REQUIRED', 'NOT_VERIFIABLE'}


def test_zero_and_positive_preserve_chemistry_and_original_gross_classification(controls):
    zero = impact(controls['main'], 0)['impact_ledger']; positive = impact(controls['main'], 10)['impact_ledger']
    assert zero['source_after_tables'] == {k: controls['dv0']['expected']['fixture_expected']['Main'][k] for k in ('blend', 'additives')}
    assert zero['source_after_total_preparation_water_full_precision_bbl'] == zero['pure_mix_water_full_precision_bbl']
    assert zero['source_invariants_and_unselected_slurries'] == positive['source_invariants_and_unselected_slurries']
    assert zero['source_after_tables']['blend'] == positive['source_after_tables']['blend']
    assert any(r['Name'] == 'C03 Dry' for r in positive['source_after_tables']['blend'])
    assert {r['material'] for r in positive['gross_additives']} == {'C03 Powder', 'C03 Liquid', 'SALT'}
    assert {r['unit'] for r in positive['gross_additives']} == {'lb', 'gal'}
    for before, after in zip(zero['gross_additives'], positive['gross_additives']):
        assert float(before['source_after_display'].removesuffix(' gal')) < float(after['source_after_display'].removesuffix(' gal'))
    z = impact(controls['zero-dose'], 0)['impact_ledger']
    assert all(float(r['source_after_display'].removesuffix(' gal')) == 0 for r in z['gross_additives'])


def test_overlapping_before_interval_is_never_an_exact_historical_delta(controls):
    same = impact(controls['main'], 5)['impact_ledger']; outside = impact(controls['main'], 10)['impact_ledger']
    interval = same['before_dv_interval']['compatible_display_dv_interval_bbl']
    assert float(interval['lower']) < 5 < float(interval['upper']) < 10
    assert same['NOTE_display_equals_original'] is True and outside['NOTE_display_equals_original'] is False
    assert same['historical_exact_delta'] is outside['historical_exact_delta'] is None
    assert same['before_dv_interval']['not_a_measured_heel'] is True


def test_original_half_up_NOTE_boundary_uses_C04_not_new_rounding(controls):
    case = controls['main']; before = inspected(str(case['path']), json.dumps(case['op'], sort_keys=True), 'Main')
    pure = before['original_c03']['c02']['reference']['slurries']['Main']['calc']['field_water_bbl']
    for proposed in (40.25-pure, 40.25-pure-0.0000001):
        result = impact(case, proposed)
        oracle = g.preview_source.preview(case['path'], case['op'], {'slurry': 'Main', 'corrected_dead_volume_bbl': proposed})
        assert result['impact_ledger']['source_after_NOTE_display_bbl'] == oracle['selected_dead_volume_preview']['new_preparation_water_note_display_bbl']
    assert impact(case, 40.25-pure)['impact_ledger']['source_after_NOTE_display_bbl'] == '40.3'


def test_lead_selection_retains_unselected_tail_and_all_C02_C03_findings(controls):
    case = controls['two']; r = impact(case, 10, 'Lead')
    before = inspected(str(case['path']), json.dumps(case['op'], sort_keys=True), 'Lead')['original_c03']
    assert r['original_c03'] == before and r['c04']['c03'] == before
    assert r['impact_ledger']['source_invariants_and_unselected_slurries']['unselected_slurry_references'] == {'Tail': before['c02']['reference']['slurries']['Tail']}
    assert r['impact_ledger']['source_invariants_and_unselected_slurries']['raw_lab_findings_sha256_before'] == r['impact_ledger']['source_invariants_and_unselected_slurries']['raw_lab_findings_sha256_after']
    edit = request(case, 10, 'Lead'); edit['slurry'] = 'Tail'
    assert g.dead_volume_impact(case['path'], case['op'], edit)['status'] == 'INPUT_REJECTED'


@pytest.mark.parametrize('value', [-1, True, '0', None, float('nan'), float('inf')])
def test_invalid_DV_fails_closed(controls, value):
    case = controls['main']; edit = request(case); edit['corrected_dead_volume_bbl'] = value
    if isinstance(value, float):
        with pytest.raises(ValueError): g.dead_volume_impact(case['path'], case['op'], edit)
    else: assert g.dead_volume_impact(case['path'], case['op'], edit)['status'] == 'INPUT_REJECTED'


@pytest.mark.parametrize('mutation', ['document', 'operator', 'source', 'template', 'claims', 'locator', 'extra', 'missing', 'duplicate'])
def test_stale_or_forged_selection_fails_closed(controls, mutation):
    case = controls['main']; edit = request(case); op = deepcopy(case['op']); path = case['path']
    if mutation == 'document':
        path = path.parent/'changed.docx'; path.write_bytes(case['path'].read_bytes())
        d = Document(path); d.add_paragraph('Synthetic later change'); d.save(path)
    elif mutation == 'operator': op['slurries']['Main']['density_pcf'] = 105
    elif mutation == 'source': edit['bindings']['source_sha'] = '0'*40
    elif mutation == 'template': edit['bindings']['template_sha256'] = '0'*64
    elif mutation == 'claims': edit['bindings']['selected_claims_sha256'] = '0'*64
    elif mutation == 'locator': edit['bindings']['forged_locator'] = {'row': 2}
    elif mutation == 'extra': edit['observed_lab'] = {'free_water': 0}
    elif mutation == 'missing': del edit['slurry']
    else: edit['slurry'] = ['Main', 'Main']
    original = path.read_bytes(); result = g.dead_volume_impact(path, op, edit)
    assert result['status'] == 'INPUT_REJECTED' and result['word_modified'] is False
    assert path.read_bytes() == original


@pytest.mark.parametrize('mutation', ['hidden-gross', 'hidden-note', 'split-gross', 'split-note', 'merged', 'missing-row',
    'malformed-gross', 'duplicate-note', 'missing-note', 'wrong-tank', 'missing-grid', 'duplicate-table', 'split-title'])
def test_unproved_physical_claims_fail_closed(controls, mutation):
    case = controls['main']; path = case['path'].parent/(mutation+'.docx'); d = Document(case['path'])
    table = next(t for t in d.tables if t.cell(0, 0).text.endswith(' Cement Slurry Additives Data'))
    note = next(p for p in d.paragraphs if 'Mix above Additives in' in p.text)
    p = note if mutation.endswith('note') else table.rows[2].cells[6].paragraphs[0]
    if mutation.startswith('hidden'): p.runs[0]._r.get_or_add_rPr().append(OxmlElement('w:vanish'))
    elif mutation in ('split-note', 'split-gross'):
        raw = p.text; p.runs[0].text = raw[:2]; p.add_run(raw[2:])
    elif mutation == 'merged': table.rows[2].cells[5].merge(table.rows[2].cells[6])
    elif mutation == 'missing-row': table._tbl.remove(table.rows[2]._tr)
    elif mutation == 'malformed-gross': p.runs[0].text = 'not a quantity'
    elif mutation == 'duplicate-note': note._p.addnext(deepcopy(note._p))
    elif mutation == 'missing-note': note._p.getparent().remove(note._p)
    elif mutation == 'wrong-tank': note.runs[0].text = note.text.replace('Mud Reserve Tanks', 'Unknown tank')
    elif mutation == 'missing-grid': table._tbl.remove(table._tbl.find(qn('w:tblGrid')))
    elif mutation == 'split-title':
        p = table.cell(0, 0).paragraphs[0]; raw = p.text; p.runs[0].text = raw[:2]; p.add_run(raw[2:])
    else: table._tbl.addnext(deepcopy(table._tbl))
    d.save(path); original = path.read_bytes()
    r = g.inspect_selection(path, case['op'], 'Main')
    assert r['status'] == 'INPUT_REJECTED' and r['word_modified'] is False and path.read_bytes() == original


@pytest.mark.parametrize('mutation', ['after-zero', 'after-oversized', 'after-malformed', 'after-negative',
    'after-missing', 'after-duplicate', 'leading', 'shifted', 'span-shift', 'width', 'header',
    'duplicate-material', 'split-material', 'inherited-hidden'])
def test_valid_trailing_grid_never_exempts_visibility_or_column_integrity(controls, mutation):
    case = controls['main']; path = case['path'].parent/('grid-'+mutation+'.docx')
    document = Document(case['path'])
    table = next(t for t in document.tables if t.cell(0, 0).text == 'Main Cement Slurry Additives Data')
    row = table._tbl.findall(qn('w:tr'))[2]; cells = row.findall(qn('w:tc'))
    props = row.find(qn('w:trPr')); after = props.find(qn('w:gridAfter'))
    assert len(table._tbl.find(qn('w:tblGrid'))) == 8 and len(cells) == 7
    assert after.get(qn('w:val')) == '1' and props.find(qn('w:gridBefore')) is None
    if mutation in ('after-zero', 'after-oversized', 'after-malformed', 'after-negative'):
        after.set(qn('w:val'), {'after-zero': '0', 'after-oversized': '9', 'after-malformed': '1.0', 'after-negative': '-1'}[mutation])
    elif mutation == 'after-missing': props.remove(after)
    elif mutation == 'after-duplicate': props.append(deepcopy(after))
    elif mutation in ('leading', 'shifted'):
        before = OxmlElement('w:gridBefore'); before.set(qn('w:val'), '1'); props.insert(0, before)
        if mutation == 'shifted': after.set(qn('w:val'), '0')
    elif mutation == 'span-shift':
        span = OxmlElement('w:gridSpan'); span.set(qn('w:val'), '2')
        cells[1].find(qn('w:tcPr')).append(span)
    elif mutation == 'width': props.find(qn('w:wAfter')).set(qn('w:w'), '9999')
    elif mutation == 'header': table.rows[1].cells[6].text = 'Different gross header'
    elif mutation == 'duplicate-material':
        other = table._tbl.findall(qn('w:tr'))[3]; other.addprevious(deepcopy(row)); table._tbl.remove(other)
    elif mutation == 'split-material':
        p = table.rows[2].cells[1].paragraphs[0]; raw = p.text; p.runs[0].text = raw[:2]; p.add_run(raw[2:])
    else:
        parent = document.styles.add_style('G06B Hidden Parent', WD_STYLE_TYPE.PARAGRAPH)
        parent.font.hidden = True
        inherited = document.styles.add_style('G06B Inherited Claim', WD_STYLE_TYPE.PARAGRAPH)
        inherited.base_style = parent; table.rows[2].cells[6].paragraphs[0].style = inherited
    document.save(path); original = path.read_bytes()
    result = g.inspect_selection(path, case['op'], 'Main')
    assert result['status'] == 'INPUT_REJECTED' and result.get('c04') is None
    assert result['docx_written'] is result['word_modified'] is False and path.read_bytes() == original


@pytest.mark.parametrize('job,slurry', [('CMT PLUG WITH CT', 'Main'), ('CMT SQUEEZE', 'Main'), ('CSG 20"', 'Scavenger'), ('CSG 20"', 'Lead')])
def test_unsupported_or_inactive_selection(controls, job, slurry):
    case = controls['main']; op = deepcopy(case['op']); op['job_type'] = job
    if slurry == 'Scavenger': op['active_fluids'] = [slurry]; op['slurries'] = {slurry: op['slurries']['Main']}
    assert g.inspect_selection(case['path'], op, slurry)['status'] == 'INPUT_REJECTED'


def test_C04_private_path_and_symlink_boundaries(controls):
    case = controls['main']; alias = case['path'].parent/'alias.docx'; alias.symlink_to(case['path'])
    for path in (alias, g.runtime.HERE/'fixtures/synthetic-main.docx', case['path'].parent/'..'/case['path'].parent.name/'main.docx'):
        assert g.inspect_selection(path, case['op'], 'Main')['status'] == 'INPUT_REJECTED'
    assert not alias.is_dir()


def test_cli_unique_JSON_and_explicit_zero(controls):
    case = controls['main']; operator = case['path'].parent/'operator.json'; edit = case['path'].parent/'edit.json'
    operator.write_text(json.dumps(case['op'])); edit.write_text(json.dumps(request(case, 0)))
    command = [sys.executable, '-B', str(g.ROOT/'g06b_dead_volume.py'), str(operator), str(case['path']), '--edit', str(edit)]
    p = subprocess.run(command, text=True, capture_output=True)
    r = json.loads(p.stdout)
    assert p.returncode == 0 and r['impact_ledger']['explicit_dead_volume_bbl'] == 0
    edit.write_text('{"slurry":"Main","slurry":"Tail"}')
    p = subprocess.run(command, text=True, capture_output=True)
    assert p.returncode == 2 and json.loads(p.stdout)['status'] == 'INPUT_REJECTED'
