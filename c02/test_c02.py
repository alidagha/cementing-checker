"""Synthetic partial comparisons, never App UAT/golden certification."""
from copy import deepcopy
from decimal import Decimal
from functools import lru_cache
import hashlib
import json
from pathlib import Path
import subprocess
import sys

from docx import Document
import pytest
import c02_compare as cc


def operator(volume=50, density=104, additives=None):
    return {'job_type':'CSG 20"','active_fluids':['Main'],'slurries':{'Main':{
        'client_volume_bbl':volume,'density_pcf':density,'base_cement':'Cement G Delijan',
        'cement_sg':3.2,'additives':[] if additives is None else additives}}}


@lru_cache(maxsize=32)
def cached_engine(text):
    return cc.engine(json.loads(text))


def original_ref(op):
    return deepcopy(cached_engine(json.dumps(op,sort_keys=True)))


@pytest.fixture(autouse=True)
def reuse_only_actual_source_results(monkeypatch):
    # Avoid repeated source imports for identical operator inputs; every cached
    # entry is from a real isolated original-source call, never a fabricated result.
    original = cc.engine
    @lru_cache(maxsize=32)
    def call(text): return original(json.loads(text))
    monkeypatch.setattr(cc, 'engine', lambda op: deepcopy(call(json.dumps(op,sort_keys=True))))


def table(doc, title, headers, rows):
    t=doc.add_table(rows=1,cols=len(headers));t.cell(0,0).merge(t.cell(0,len(headers)-1)).text=title
    for values in [headers,*rows]:
        cells=t.add_row().cells
        for c,v in zip(cells,values):c.text=str(v)
    return t


def fixture(path, values=None, fluid='Main', duplicate=False, labrows=None, addrows=None):
    v={'yield':'1.818','sacks':'154.4','water':'999','density':'104'}
    v.update(values or {})
    d=Document();d.add_paragraph('Executive Summary');d.add_paragraph('Geometric volume 9999 bbl is INFO only.')
    d.add_paragraph('V. Cementing Program')
    for _ in range(2 if duplicate else 1):
        # Program data has paired labels; no invented pure-water alias.
        t=d.add_table(rows=1,cols=4);t.cell(0,0).merge(t.cell(0,3)).text=fluid+' Cement Slurry Data'
        for row in [('Yield',v['yield']+' cuft/sk','Slurry Weight',v['density']+' pcf'),
                    ('Solution',v['water']+' bbl','Volume','50 bbl')]:
            for c,x in zip(t.add_row().cells,row):c.text=x
        table(d,fluid+' Cement Slurry Blend Data',['Item','Name','Material Type','sks or lbs','lbs / sk'],
              [['1','Cement G Delijan','Cement',v['sacks']+' sks','110.0']])
    if addrows:
        table(d,fluid+' Cement Slurry Additives Data',['Item','Name','Material Type','lbs or gal','(lbs or gal)/sk','(lbs or gal)/bbl','lbs or gal (with dead Vol.)'],addrows)
    d.add_paragraph('NOTE 7: Mix above Additives in 999 bbl Fresh Water at Mud Reserve Tanks.')
    d.add_paragraph('VII. Laboratory Report');d.add_paragraph(fluid+' Cement Slurry Results')
    table(d,'Conventional',['Material Type','Code','Lot Number','Concentration','Design Unit','Mass (g/600 mL)'],
          labrows if labrows is not None else [['Cement','CEMENT G DELIJAN','-','100.0%','BWOB','582.0'],['Base Water','Base Water','-','1.0','BWOW','418.3']])
    d.save(path);return path


def mapped(result, field):return [x for x in result['findings'] if x['field']==field]


def test_negative_controls_wrong_numbers_not_pass(tmp_path):
    p=fixture(tmp_path/'wrong.docx',{'yield':'9.9','sacks':'999','water':'999'})
    r=cc.compare_document(p,operator())
    assert {x['result'] for f in ('yield','sacks') for x in mapped(r,f)}=={'DISCREPANCY'}
    assert not any(x['result']=='MATCH' for x in r['findings'] if '999 bbl' in x['raw_reported_text'])
    assert r['status']=='PARTIAL_FINDINGS_ONLY'


def test_negative_controls_cross_slurry_duplicates_and_invalid_inputs(tmp_path):
    for p in (fixture(tmp_path/'cross.docx',fluid='Tail'),fixture(tmp_path/'dup.docx',duplicate=True)):
        r=cc.compare_document(p,operator())
        assert not any(x['result']=='MATCH' for x in r['findings'] if x['field'] in ('yield','sacks') or x['slurry']=='Tail')
    bad=operator();bad['slurries']['Main']['density_pcf']=0
    assert cc.compare_document(fixture(tmp_path/'bad.docx'),bad)['status']=='REFERENCE_INPUT_ERROR'


def test_independently_pinned_neat_reference_and_source_calls():
    r=original_ref(operator());s=r['slurries']['Main']
    # Existing source guardrail invariants: 104 pcf / SG3.2 neat → 1.818 and 582.0g.
    assert s['display']['yield_cuft_sk']=='1.818'
    assert s['lab'][0]['Mass']=='582.0'
    assert abs(s['calc']['yield_ft3_per_sk']-1.8179086538461537)<1e-12
    assert {'phase_5_cement.py:build_components','phase_5_cement.py:calculate_base_results',
            'phase_5_cement.py:build_cement_tables','phase_7_lab.py:build_lab_df_from_phase5',
            'engineering_tools.py:calculate_slurry_from_components','phase_10_procedure.py:_word_program_quantities'} <= set(r['source_calls'])
    assert not any('session_state' in c for c in r['source_calls'])


def test_correct_document_and_cli_preserve_original(tmp_path):
    p=fixture(tmp_path/'correct.docx');before=p.read_bytes()
    r=cc.compare_document(p,operator())
    assert [x['result'] for x in mapped(r,'yield')+mapped(r,'sacks')+mapped(r,'density')]==['MATCH']*3
    assert len(mapped(r,'lab_mass'))==2 and all(x['result']=='MATCH' for x in mapped(r,'lab_mass'))
    assert any(x['result']=='UNIMPLEMENTED' for x in r['findings'])
    op=tmp_path/'operator.json';op.write_text(json.dumps(operator()))
    proc=subprocess.run([sys.executable,'-B',str(cc.HERE/'c02_compare.py'),str(op),str(p)],capture_output=True,text=True)
    assert proc.returncode==0,proc.stderr
    assert json.loads(proc.stdout)['status']=='PARTIAL_FINDINGS_ONLY'
    assert p.read_bytes()==before and r['source_document_sha256']==hashlib.sha256(before).hexdigest()


@pytest.mark.parametrize('field,expected,tolerance,unit',[('yield','1.818','.001','cuft/sk'),('sacks','154.4','1','sks'),('pure_water','34.8','1','bbl'),('lab','582.0','1','')])
def test_numeric_policy_inclusive_boundaries(field,expected,tolerance,unit):
    e,t=Decimal(expected),Decimal(tolerance)
    for sign in (-1,1):
        raw=str(e+sign*t)+(' '+unit if unit else '')
        assert cc.assess(raw,unit,expected,tolerance)['result']=='MATCH'
        raw=str(e+sign*(t+Decimal('.00001')))+(' '+unit if unit else '')
        assert cc.assess(raw,unit,expected,tolerance)['result']=='DISCREPANCY'
    if field=='pure_water':
        for difference in ('.5','.75','1'):
            result=cc.assess(str(e+Decimal(difference))+' bbl',unit,expected,tolerance)
            assert result['result']=='MATCH' and Decimal(result['delta'])==Decimal(difference)


def test_wrong_density_exact_and_no_unit_coercion(tmp_path):
    r=cc.compare_document(fixture(tmp_path/'density.docx',{'density':'104.01'}),operator())
    assert mapped(r,'density')[0]['result']=='DISCREPANCY'
    assert cc.assess('1.818 bbl','cuft/sk','1.818','.001')['result']=='NOT_VERIFIABLE'


@pytest.mark.parametrize('raw,state',[('', 'MISSING'),('0','ZERO'),('-','DASH'),('N/A','NA_PRINTED'),('abc','MALFORMED_OR_UNIT_UNPROVEN')])
def test_states_are_distinct(raw,state):
    r=cc.assess(raw,'','1','0')
    assert r['reported_state']==state
    assert r['result']!='MATCH'
    if state in ('MISSING','DASH','NA_PRINTED'):
        assert cc.assess(raw+' pcf','pcf','104','0')['reported_state']==state


def custom(dose,state='Powder'):
    return {'Material Type':'Custom '+state,'Name':'Explicit Custom '+state,'Physical State':state,
            'Mix Method':'Dry Blend' if state=='Liquid' else 'In Mix Water','User Input':dose,'Density':1.4}


@pytest.mark.parametrize('dose,state,tol',[(.009,'Powder','0'),(.01,'Powder','.3'),(1,'Powder','.3'),(1.001,'Powder','1'),(.2,'Liquid','.3')])
def test_lab_additive_policy_mapping_edges(tmp_path,dose,state,tol):
    a=custom(dose,state);op=operator(additives=[a]);ref=original_ref(op)['slurries']['Main']
    source=next(x for x in ref['lab'] if x['Code']==a['Name'])
    expected=Decimal(source['Mass']);t=Decimal(tol)
    for delta,verdict in [(t,'MATCH'),(t+Decimal('.001'),'DISCREPANCY')]:
        row=[source['Material Type'],source['Code'],'-',source['Concentration'],source['Design Unit'],str(expected+delta)]
        p=fixture(tmp_path/('add'+str(delta)+'.docx'),labrows=[row])
        r=cc.compare_document(p,op);finding=mapped(r,'lab_mass')[0]
        assert finding['tolerance']==tol and finding['result']==verdict
        assert finding['full_precision_expected']==ref['raw_lab'][a['Name'].casefold()]


def test_volume_only_change_and_full_precision_lab():
    a=original_ref(operator(50))['slurries']['Main'];b=original_ref(operator(100))['slurries']['Main']
    assert b['calc']['field_sacks']==2*a['calc']['field_sacks']
    assert b['calc']['field_water_bbl']==2*a['calc']['field_water_bbl']
    assert b['calc']['yield_ft3_per_sk']==a['calc']['yield_ft3_per_sk']
    assert b['calc']['lab_cmt_gr']==a['calc']['lab_cmt_gr'] and b['calc']['lab_water_gr']==a['calc']['lab_water_gr']
    assert a['calc']['lab_cmt_gr']!=1058/float(a['display']['yield_cuft_sk'])


def test_source_additive_classification_and_invalid_identity(tmp_path):
    a=custom(.1,'Liquid');r=original_ref(operator(additives=[a]))['slurries']['Main']
    assert r['calc']['field_liquids'] and not r['calc']['field_powders']
    for change in ({'Name':'Other (Custom)'},{'Density':None},{'User Input':-1}):
        row=custom(.1);row.update(change)
        got=cc.compare_document(fixture(tmp_path/('bad'+str(len(str(change)))+'.docx')),operator(additives=[row]))
        assert got['status']=='REFERENCE_INPUT_ERROR'


def test_two_positive_slurry_identities(tmp_path):
    op=operator();op['active_fluids']=['Lead','Tail'];main=op['slurries'].pop('Main')
    op['slurries']={'Lead':main,'Tail':dict(main,density_pcf=118)}
    for fluid,yield_text in [('Lead','1.818'),('Tail','1.360')]:
        r=cc.compare_document(fixture(tmp_path/(fluid+'.docx'),{'yield':yield_text},fluid=fluid),op)
        assert mapped(r,'yield')[0]['slurry']==fluid and mapped(r,'yield')[0]['result']=='MATCH'


def test_no_required_optional_fluids_or_original_dead_volume(tmp_path):
    p=fixture(tmp_path/'note.docx')
    r=cc.compare_document(p,operator())
    assert 'Spacer' not in r['reference']['slurries']
    assert all(x['result']!='MATCH' for x in r['findings'] if 'Mix above' in x['raw_reported_text'])
    assert not any(x['field']=='pure_water' for x in r['findings'])


def test_field_additive_source_display_exact_and_duplicate_fail_closed(tmp_path):
    a=custom(.1);op=operator(additives=[a]);s=original_ref(op)['slurries']['Main'];row=s['additives'][0]
    data=[[str(row[k]) for k in ['Item','Name','Material Type','lbs or gal','(lbs or gal)/sk','(lbs or gal)/bbl','lbs or gal (with dead Vol.)']]]
    p=fixture(tmp_path/'field.docx',addrows=data)
    assert mapped(cc.compare_document(p,op),'field_additive')[0]['result']=='MATCH'
    wrong=deepcopy(data);wrong[0][3]=str(Decimal(wrong[0][3])+Decimal('.1'))
    assert mapped(cc.compare_document(fixture(tmp_path/'wrongfield.docx',addrows=wrong),op),'field_additive')[0]['result']=='DISCREPANCY'
    dup=cc.compare_document(fixture(tmp_path/'dupfield.docx',addrows=data+data),op)
    assert all(r['result']=='NOT_VERIFIABLE' for r in mapped(dup,'field_additive'))
    assert all(r['mapping_state']=='AMBIGUOUS' and r['association_state']=='EXPLICIT' for r in mapped(dup,'field_additive'))


def test_operator_volume_not_docx_or_geometry_controls_sacks(tmp_path):
    r=cc.compare_document(fixture(tmp_path/'volume.docx'),operator(100))
    f=mapped(r,'sacks')[0]
    assert f['source_rounded_expected']=='308.8' and f['result']=='DISCREPANCY'
    assert mapped(r,'yield')[0]['result']=='MATCH'


def test_duplicate_json_and_missing_explicit_sg_rejected(tmp_path):
    from c02_engine import read_json
    with pytest.raises(ValueError):read_json('{"job_type":"a","job_type":"b"}')
    op=operator();del op['slurries']['Main']['cement_sg']
    assert cc.compare_document(fixture(tmp_path/'missing-sg.docx'),op)['status']=='REFERENCE_INPUT_ERROR'


def test_scavenger_does_not_invent_placed_or_independent_lab_report(tmp_path):
    op=operator();op['active_fluids']=['Scavenger'];op['slurries']={'Scavenger':op['slurries']['Main']}
    r=cc.compare_document(fixture(tmp_path/'scavenger.docx',fluid='Scavenger'),op)
    assert not mapped(r,'lab_mass')
    assert not any(x['result']=='MATCH' for x in r['findings'])


def test_two_slurries_in_one_docx_and_crossed_values(tmp_path):
    op=operator();main=op['slurries'].pop('Main');op['active_fluids']=['Lead','Tail']
    op['slurries']={'Lead':main,'Tail':dict(main,density_pcf=118)}
    for crossed in (False,True):
        doc=Document();doc.add_paragraph('V. Cementing Program')
        for fluid,value in [('Lead','1.360' if crossed else '1.818'),('Tail','1.818' if crossed else '1.360')]:
            t=doc.add_table(rows=1,cols=2);t.cell(0,0).merge(t.cell(0,1)).text=fluid+' Cement Slurry Data'
            for c,text in zip(t.add_row().cells,['Yield',value+' cuft/sk']):c.text=text
        path=tmp_path/('two-crossed.docx' if crossed else 'two-correct.docx');doc.save(path)
        findings=mapped(cc.compare_document(path,op),'yield')
        assert [x['slurry'] for x in findings]==['Lead','Tail']
        assert [x['result'] for x in findings]==['DISCREPANCY' if crossed else 'MATCH']*2
