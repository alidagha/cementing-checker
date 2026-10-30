"""C03 source-rendered SYNTHETIC structural controls; no goldens/UAT."""
from copy import deepcopy
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import subprocess
import sys

from docx import Document
import pytest
import c03_water as cw
from c03_source import HERE,PIN,request


def operator(mixed=False,two=False):
    op=json.loads((HERE.parent/'c02/fixtures/operator-neat.json').read_text())
    if mixed:
        def row(name,kind,state,mix,dose,density):
            return {'Name':name,'Material Type':kind,'Physical State':state,'Mix Method':mix,'User Input':dose,'Density':density}
        op['slurries']['Main']['additives']=[
            row('C03 Powder','Custom Powder','Powder','In Mix Water',1,1.4),
            row('C03 Liquid','Custom Liquid','Liquid','Dry Blend',.1,1.2),
            row('C03 Dry','Custom Powder','Powder','Dry Blend',2,2.5),
            row('SALT','Nacl','Powder','In Mix Water',5,2.16)]
    if two:
        main=op['slurries'].pop('Main');op['active_fluids']=['Lead','Tail']
        op['slurries']={'Lead':main,'Tail':dict(main,density_pcf=118)}
        op['job_type']='CSG 9 5/8"'
    return op


@pytest.fixture(scope='session')
def cases(tmp_path_factory):
    directory=tmp_path_factory.mktemp('rendered')
    cases={}
    for name,mixed,two,dead in [('neat0',False,False,0),('neat5',False,False,5),('mixed0',True,False,0),('mixed5',True,False,5),('two',False,True,5),('zero-dose',True,False,5)]:
        op=operator(mixed,two);path=directory/(name+'.docx')
        if name=='zero-dose':
            for row in op['slurries']['Main']['additives']:row['User Input']=0
        setup={fluid:{'dead':dead if fluid!='Tail' else 2,'tank':'Batch Mixer Tanks' if fluid=='Tail' else 'Mud Reserve Tanks'} for fluid in op['active_fluids']}
        rendered=request({'action':'render_fixture','operator':op,'docx':str(path),'fixture':setup})
        inspected=request({'action':'inspect','operator':op,'docx':str(path)})
        result=cw.extend(inspected,op)
        cases[name]={'path':path,'operator':op,'rendered':rendered,'inspected':inspected,'result':result}
        # Durable copies and fixture metadata (pytest temp remains confined).
        (HERE/'fixtures'/(name+'.docx')).write_bytes(path.read_bytes())
        (HERE/'fixtures'/(name+'-operator.json')).write_text(json.dumps(op,indent=2))
        (HERE/'fixtures'/(name+'-source.json')).write_text(json.dumps(rendered,indent=2))
        (HERE/'output'/(name+'.json')).write_text(json.dumps(result,indent=2))
    return cases


def mutate(case,tmp_path,change):
    document=Document(case['path']);change(document)
    path=tmp_path/'mutated.docx';document.save(path)
    before=path.read_bytes();result=cw.analyze(path,case['operator'])
    assert path.read_bytes()==before
    return result


def notes(doc):return [p for p in doc.paragraphs if 'Mix above Additives in' in p.text]
def water(result,fluid='Main'):return next(f for f in result['water_findings'] if f['slurry']==fluid)
def adds(doc):return next(t for t in doc.tables if t.cell(0,0).text=='Main Cement Slurry Additives Data')


def test_genuine_red_cases_green_and_baseline_preserved(cases):
    case=cases['neat5'];result=case['result'];w=water(result)
    assert w['status']=='INFERRED_AT_DISPLAY_PRECISION'
    assert w['Mixing Tank Selected in Word']=='Mud Reserve Tanks'
    # Independent pinned neat reference, not a projection compared with itself.
    assert abs(w['Pure Mix Water Required']['full_precision_bbl']-34.848484848484844)<1e-12
    assert w['Total Preparation Water Reported in Word']['raw']=='39.8'
    inferred=w['Dead Volume Included by Engineer'];interval=inferred['compatible_display_dv_interval_bbl']
    assert Decimal(interval['lower'])<5<Decimal(interval['upper'])
    assert Decimal(inferred['nominal_bbl'])!=5 and inferred['not_a_measured_heel']
    assert w['locator']['body_block'] and len(w['proof_table_locators'])==3 and w['note']
    assert result['c02']==case['inspected']['c02']
    assert result['status']=='PARTIAL_FINDINGS_ONLY'
    assert any(f['result']=='UNIMPLEMENTED' for f in result['c02']['findings'])
    assert not any(f['field']=='pure_water' and f['result']=='MATCH' for f in result['c02']['findings'])
    # The actual pre-implementation fixture and captured baseline stay intact.
    old=json.loads((HERE/'evidence/red-baseline.json').read_text())
    assert water(cw.extend(old,operator()))['status']=='INFERRED_AT_DISPLAY_PRECISION'


def test_zero_allowance_and_negative_nominal_not_silently_clamped(cases):
    inferred=water(cases['neat0']['result'])['Dead Volume Included by Engineer']
    assert Decimal(inferred['nominal_bbl'])<0
    assert inferred['nominal_source_valid'] is False
    assert Decimal(inferred['source_feasible_interval_bbl']['lower'])==0
    assert inferred['samples']['lower_enclosure']==0
    assert inferred['qualification']=='INFERRED_AT_DISPLAY_PRECISION'


def test_two_different_tanks_and_no_first_slurry_borrow(cases):
    r=cases['two']['result'];lead,tail=water(r,'Lead'),water(r,'Tail')
    assert lead['Mixing Tank Selected in Word']=='Mud Reserve Tanks'
    assert tail['Mixing Tank Selected in Word']=='Batch Mixer Tanks'
    for fluid,actual in [('Lead',5),('Tail',2)]:
        w=water(r,fluid);assert w['status']=='INFERRED_AT_DISPLAY_PRECISION'
        i=w['Dead Volume Included by Engineer']['compatible_display_dv_interval_bbl']
        assert Decimal(i['lower'])<=actual<Decimal(i['upper'])
    assert lead['locator']!=tail['locator']


def test_source_gross_powder_liquid_salt_and_dry_blend(cases):
    a,b=cases['mixed0']['result'],cases['mixed5']['result']
    for r in (a,b):
        assert r['source_projections']['base_invariants_preserved']
        findings=r['gross_additive_findings'];assert {f['material'] for f in findings}=={'C03 Powder','C03 Liquid','SALT'}
        assert all(f['status'] in ('INFERRED_AT_DISPLAY_PRECISION','MATCH_AT_DISPLAY_PRECISION') for f in findings)
        assert {f['unit'] for f in findings}=={'lb','gal'}
        assert all(f['source_function']=='phase_5_cement.build_cement_tables' for f in findings)
    for f in a['gross_additive_findings']:
        g=next(x for x in b['gross_additive_findings'] if x['material']==f['material'])
        assert Decimal(f['raw_text'].replace(' gal',''))<Decimal(g['raw_text'].replace(' gal',''))
    sa=a['c02']['reference']['slurries']['Main'];sb=b['c02']['reference']['slurries']['Main']
    assert sa['calc']==sb['calc'] and sa['lab']==sb['lab'] and sa['blend']==sb['blend']
    assert any(x['Name']=='C03 Dry' for x in sa['blend'])
    assert not any(x['material']=='C03 Dry' for x in b['gross_additive_findings'])


@pytest.mark.parametrize('kind',['missing','duplicate','misplaced','contradictory_tank','unknown_tank','aggregate_only','solution_decoy'])
def test_missing_ambiguous_or_unproved_water_no_positive(cases,tmp_path,kind):
    def change(d):
        n=notes(d)[0]
        if kind=='missing':n._element.getparent().remove(n._element)
        elif kind=='duplicate':n._element.addnext(deepcopy(n._element))
        elif kind=='misplaced':d.paragraphs[-1]._element.addnext(n._element)
        elif kind=='contradictory_tank':
            target=next(p for p in d.paragraphs if p.text=='X. Cementing Procedure')
            total=water(cases['neat5']['result'])['Total Preparation Water Reported in Word']['raw']
            p=deepcopy(n._element);target._element.addnext(p)
            from docx.text.paragraph import Paragraph
            Paragraph(p,target._parent).text=f'1. Fill up Batch Mixer Tanks with Fresh water (Main: {total} bbl) for mixing fluids.'
        elif kind=='unknown_tank':n.text=n.text.replace('Mud Reserve Tanks','Tank Guess')
        elif kind=='aggregate_only':
            n.text=''
            target=next(p for p in d.paragraphs if p.text=='X. Cementing Procedure')
            p=deepcopy(n._element);target._element.addnext(p)
            from docx.text.paragraph import Paragraph
            Paragraph(p,target._parent).text='1. Fill up Mud Reserve Tanks with Fresh water (Main: 39.8 bbl & Tail: 15.0 bbl) for mixing fluids.'
        elif kind=='solution_decoy':n.text='NOTE 5: Solution 39.8 bbl at Mud Reserve Tanks.'
    r=mutate(cases['neat5'],tmp_path,change)
    assert water(r)['status']=='NOT_VERIFIABLE'
    assert not r['source_projections']
    if kind=='aggregate_only':
        assert r['procedure_water_evidence'] and r['procedure_water_evidence'][0]['fluid'] is None


@pytest.mark.parametrize('token,state',[(' bbl','MISSING'),('0.0 bbl','ZERO'),('- bbl','DASH'),('N/A bbl','NA_PRINTED'),('oops bbl','MALFORMED_OR_UNPROVEN_PRECISION'),('-1.0 bbl','NEGATIVE'),('39.8 m3','NUMBER'),('1.0 bbl','NUMBER'),('39.80 bbl','MALFORMED_OR_UNPROVEN_PRECISION')])
def test_distinct_states_units_and_inconsistent_total(cases,tmp_path,token,state):
    def change(d):notes(d)[0].text=f'NOTE 5: Mix above Additives in {token} Fresh Water at Mud Reserve Tanks.'
    w=water(mutate(cases['neat5'],tmp_path,change))
    assert w['status']=='NOT_VERIFIABLE'
    assert w['Total Preparation Water Reported in Word']['state']==state
    if token in ('0.0 bbl','1.0 bbl','-1.0 bbl'):
        assert w['Dead Volume Included by Engineer']['qualification']=='INCONSISTENT_WITH_SOURCE_NONNEGATIVE_ALLOWANCE'


def test_wrong_water_retains_gross_discrepancy_control(cases,tmp_path):
    def change(d):notes(d)[0].text=notes(d)[0].text.replace(' in ',' in 9',1)
    r=mutate(cases['mixed5'],tmp_path,change)
    # A larger Word total is a claimed larger allowance, not independently a
    # physical-water FAIL. Its inconsistent printed gross rows are discrepancies.
    assert water(r)['status']=='INFERRED_AT_DISPLAY_PRECISION'
    assert all(f['status']=='DISCREPANCY' for f in r['gross_additive_findings'])


@pytest.mark.parametrize('kind',['wrong','duplicate','wrong_material_column','wrong_unit'])
def test_gross_negative_controls(cases,tmp_path,kind):
    def change(d):
        t=adds(d)
        if kind=='wrong':t.rows[2].cells[6].text='99999.9'
        elif kind=='duplicate':t._tbl.append(deepcopy(t.rows[2]._tr))
        elif kind=='wrong_material_column':t.rows[1].cells[1].text='Unproved material'
        elif kind=='wrong_unit':t.rows[2].cells[6].text='193.6 kg'
    r=mutate(cases['mixed5'],tmp_path,change)
    f=next(x for x in r['gross_additive_findings'] if x['material']=='C03 Powder')
    if kind in ('wrong','wrong_unit'):assert f['raw_text']==('99999.9' if kind=='wrong' else '193.6 kg')
    assert f['status']==('DISCREPANCY' if kind=='wrong' else 'NOT_VERIFIABLE')


def test_same_tank_aggregate_ambiguity_and_missing_second_note(cases,tmp_path):
    def change(d):
        nn=notes(d);nn[1].text=nn[0].text
    r=mutate(cases['two'],tmp_path,change)
    assert all(w['status']=='NOT_VERIFIABLE' for w in r['water_findings'])
    assert not r['source_projections']


def test_source_sha_drift_and_document_identity_mismatch(cases):
    for field in ('source_sha','document'):
        data=deepcopy(cases['neat5']['inspected'])
        if field=='source_sha':data['source_sha']='0'*40
        else:data['parsed']['source_document_sha256']='0'*64
        with pytest.raises(ValueError):cw.extend(data,cases['neat5']['operator'])
    op=deepcopy(cases['neat5']['operator']);op['slurries']['Main']['client_volume_bbl']=99
    with pytest.raises(ValueError):cw.extend(cases['neat5']['inspected'],op)


def test_original_docx_unchanged_deterministic_cli_json(cases,tmp_path):
    case=cases['mixed5'];before=case['path'].read_bytes()
    path=tmp_path/'operator.json';path.write_text(json.dumps(case['operator']))
    p=subprocess.run([sys.executable,'-B',str(HERE/'c03_water.py'),str(path),str(case['path'])],capture_output=True,text=True)
    assert p.returncode==0,p.stderr
    r=json.loads(p.stdout);assert r==case['result'] and r['status']=='PARTIAL_FINDINGS_ONLY'
    assert case['path'].read_bytes()==before
    assert r['source_document_sha256']==hashlib.sha256(before).hexdigest()


def test_no_network_or_write_escape_and_input_symlink(cases,tmp_path):
    code="""from c03_source import install_audit
from pathlib import Path
import socket
install_audit()
for action in (lambda: Path('/tmp/c03-disallowed-write').write_text('x'),lambda: socket.getaddrinfo('example.com',443)):
    try: action()
    except RuntimeError: pass
    else: raise AssertionError('audit failed')
print('BLOCKED writes/network')
"""
    p=subprocess.run([sys.executable,'-B','-c',code],capture_output=True,text=True,cwd=HERE)
    assert p.returncode==0 and 'BLOCKED' in p.stdout,p.stderr
    target=HERE.parent/'c02/fixtures/independent-neat-claims.docx';assert target.exists()
    link=tmp_path/'escape.docx';link.symlink_to(target)
    with pytest.raises(ValueError):cw.analyze(link,cases['neat5']['operator'])
    with pytest.raises(ValueError):request({'action':'render_fixture','operator':operator(),'fixture':{'Main':{'dead':0,'tank':'Mud Reserve Tanks'}},'docx':'/tmp/escape.docx'})


def test_rounding_enclosure_boundaries_and_no_false_exactness():
    i=cw.infer('35.0',34.95)
    assert i['nominal_bbl']=='0.05'
    assert i['compatible_display_dv_interval_bbl']['lower']=='0.00'
    assert i['compatible_display_dv_interval_bbl']['upper']=='0.10'
    assert i['samples']['lower_enclosure']==0 and i['samples']['upper_enclosure']>.1
    assert i['qualification']=='INFERRED_AT_DISPLAY_PRECISION'


def test_rejected_negative_half_up_endpoint_annotation():
    i=cw.infer('-1.0',34.95)
    assert i['samples'] is None and i['qualification']=='INCONSISTENT_WITH_SOURCE_NONNEGATIVE_ALLOWANCE'
    assert i['compatible_display_dv_interval_bbl']['lower_inclusive'] is False
    assert i['compatible_display_dv_interval_bbl']['upper_inclusive'] is True


def test_invariant_source_display_match_only_for_zero_dose(cases):
    findings=cases['zero-dose']['result']['gross_additive_findings']
    assert len(findings)==3 and all(f['status']=='MATCH_AT_DISPLAY_PRECISION' for f in findings)
    assert all(f['conservative_source_display_range']==['0.0','0.0'] for f in findings)
    assert all(set(f['delta_to_source_display_projections'].values())=={'0.0'} for f in findings)


@pytest.mark.parametrize('outside',[False,True])
def test_source_display_enclosure_edge_and_outside(cases,tmp_path,outside):
    f=cases['mixed5']['result']['gross_additive_findings'][0]
    lo=Decimal(f['conservative_source_display_range'][0]);value=lo-Decimal('.1') if outside else lo
    r=mutate(cases['mixed5'],tmp_path,lambda d:setattr(adds(d).rows[2].cells[6],'text',str(value)))
    g=next(g for g in r['gross_additive_findings'] if g['material']=='C03 Powder')
    assert g['raw_text']==str(value)
    assert g['status']==('DISCREPANCY' if outside else 'INFERRED_AT_DISPLAY_PRECISION')


def test_preserved_gross_red_fixture_now_attributed():
    op=operator();op['slurries']['Main']['additives']=[{'Material Type':'Custom Powder','Name':'C03 Powder',
        'Physical State':'Powder','Mix Method':'In Mix Water','User Input':1,'Density':1.4}]
    before=(HERE/'fixtures/red-gross.docx').read_bytes()
    r=cw.analyze(HERE/'fixtures/red-gross.docx',op)
    g=r['gross_additive_findings'][0]
    assert g['raw_text']=='193.6' and g['material']=='C03 Powder'
    assert g['status']=='INFERRED_AT_DISPLAY_PRECISION'
    assert (HERE/'fixtures/red-gross.docx').read_bytes()==before


def test_missing_second_note_never_borrows_first(cases,tmp_path):
    def change(d):
        n=notes(d)[1];n._element.getparent().remove(n._element)
    r=mutate(cases['two'],tmp_path,change)
    assert water(r,'Lead')['status']=='INFERRED_AT_DISPLAY_PRECISION'
    assert water(r,'Tail')['status']=='NOT_VERIFIABLE'
    assert water(r,'Tail')['Mixing Tank Selected in Word'] is None


def test_split_runs_preserve_exact_note(cases,tmp_path):
    def change(d):
        n=notes(d)[0];text=n.text;n.text=''
        for part in (text[:20],text[20:40],text[40:]):n.add_run(part)
    r=mutate(cases['neat5'],tmp_path,change)
    assert water(r)['raw_text']==water(cases['neat5']['result'])['raw_text']
    assert water(r)['status']=='INFERRED_AT_DISPLAY_PRECISION'
